"""
Word (.docx) report and figure helpers shared by the Mayim Tools reports
(Hydrological Tools, Soil Tools) - zero QGIS dependency. python-docx and
matplotlib are imported lazily inside the functions, so tools still work
without them. Figures
use matplotlib's object API (no pyplot), which is safe on QGIS's
Processing worker thread.

Originally in hydrology/_common/ (moved here so every tool group can use
it); hydrology/_common/docx_report.py re-exports these names so existing
imports keep working.
"""

from __future__ import annotations

import io

# Chart styling - a validated categorical palette, recessive grid and ink.
C_BLUE = "#2a78d6"
C_ORANGE = "#eb6834"
C_AQUA = "#1baf7a"
C_INK = "#0b0b0b"
C_INK_2 = "#52514e"
C_GRID = "#e4e3df"
FIG_WIDTH_CM = 16.0
HOUSE_FIG_WIDTH_CM = 16.5  # Mayim house figure width


def format_value(value, fmt, scale) -> str:
    if value is None or value == "":
        return "-"
    if scale is not None:
        value = value * scale
    try:
        return fmt.format(value)
    except (ValueError, TypeError):
        return str(value)


def new_figure(height_cm: float):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=(FIG_WIDTH_CM / 2.54, height_cm / 2.54), dpi=200)
    FigureCanvasAgg(fig)
    return fig


def style_axes(ax):
    ax.grid(True, color=C_GRID, linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(C_INK_2)
        ax.spines[side].set_linewidth(0.6)
    ax.tick_params(colors=C_INK_2, labelsize=7, width=0.6)
    ax.xaxis.label.set_color(C_INK)
    ax.yaxis.label.set_color(C_INK)
    ax.xaxis.label.set_size(8)
    ax.yaxis.label.set_size(8)


def png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    buf.seek(0)
    return buf


# Mayim house text style (opt-in with Doc(house_style=True)):
#   normal text   Arial 10 pt black
#   headings      Arial 12 pt bold, #0F0F36, all caps   (heading levels 1-2)
#   sub-headings  Arial 10 pt bold, #0F0F36             (heading level 3+, labels)
#   captions      Arial 9 pt italic (table and figure headings)
HOUSE_FONT = "Arial"
HOUSE_HEADING_RGB = "0F0F36"
HOUSE_TEXT_RGB = "000000"
HOUSE_TABLE_HEADER = "ADADCC"  # table header row shading
HOUSE_TITLE_SIZE = 20  # title page title, pt
TITLE_PAGE_TOP_CM = 7.0  # space above the title (about a third down the page)
PREPARED_BY = "Mayim Consulting"


def _set_style_font(style, name, size_pt, bold=None, italic=None, rgb=None, caps=None):
    """Sets a style's font explicitly for every script (ascii, hAnsi,
    east-Asian, complex) and drops theme-font links, so Word shows exactly
    this font rather than the theme's."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    font = style.font
    font.name = name
    font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if caps is not None:
        font.all_caps = caps
    if rgb is not None:
        font.color.rgb = RGBColor.from_string(rgb)
        color = style.element.rPr.find(qn("w:color"))
        if color is not None:
            for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                color.attrib.pop(qn(attr), None)
    r_fonts = style.element.rPr.find(qn("w:rFonts"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        r_fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), name)


class Doc:
    """Thin wrapper keeping table/figure numbering. With house_style=True
    the Mayim house text style above is applied and figures are placed at
    the house figure width (16.5 cm)."""

    def __init__(self, house_style: bool = False):
        from docx import Document
        from docx.shared import Pt

        self.doc = Document()
        self.house_style = house_style
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(10)
        if house_style:
            self._apply_house_style()
        self.n_table = 0
        self.n_figure = 0
        # body tables: "center" (default) or "left" (left border on the
        # text margin)
        self.table_align = "center"

    def _apply_house_style(self):
        styles = self.doc.styles
        _set_style_font(styles["Normal"], HOUSE_FONT, 10, rgb=HOUSE_TEXT_RGB)
        for level in (1, 2):
            _set_style_font(
                styles[f"Heading {level}"],
                HOUSE_FONT,
                12,
                bold=True,
                italic=False,
                rgb=HOUSE_HEADING_RGB,
                caps=True,
            )
        for level in (3, 4):
            _set_style_font(
                styles[f"Heading {level}"],
                HOUSE_FONT,
                10,
                bold=True,
                italic=False,
                rgb=HOUSE_HEADING_RGB,
                caps=False,
            )
        caption = styles["Caption"]
        _set_style_font(caption, HOUSE_FONT, 9, bold=False, italic=True)
        caption.font.color.rgb = None

    def title_page(self, title, purpose=None, details=(), subtitle=None):
        """Title page (report standard 1.4) followed by a page break: the
        title (Arial 20 pt bold, #0F0F36, centred), an optional subtitle,
        a short purpose and goal statement and a borderless table of report
        details [(label, value)] - missing values are shown as "-"."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor

        style = self.doc.styles["Title"]
        _set_style_font(
            style,
            HOUSE_FONT,
            HOUSE_TITLE_SIZE,
            bold=True,
            italic=False,
            rgb=HOUSE_HEADING_RGB,
            caps=False,
        )
        p_pr = style.element.get_or_add_pPr()
        for border in p_pr.findall(qn("w:pBdr")):  # template's rule below
            p_pr.remove(border)
        r_pr = style.element.get_or_add_rPr()
        for tag in ("w:spacing", "w:kern"):
            for el in r_pr.findall(qn(tag)):
                r_pr.remove(el)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = self.doc.add_paragraph(title, style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(TITLE_PAGE_TOP_CM)
        p.paragraph_format.space_after = Pt(12)

        if subtitle:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(subtitle)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string(HOUSE_HEADING_RGB)
            p.paragraph_format.space_after = Pt(12)

        if purpose:
            p = self.doc.add_paragraph(purpose)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(24)

        details = list(details)
        if details:
            t = self.doc.add_table(rows=0, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            tbl_pr = t._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "nil")
                borders.append(el)
            tbl_pr.append(borders)
            for label, value in details:
                text = "-" if value is None or str(value).strip() == "" else value
                cells = t.add_row().cells
                for cell, content, bold in (
                    (cells[0], label, True),
                    (cells[1], str(text), False),
                ):
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(str(content))
                    run.bold = bold
                    run.font.size = Pt(10)
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tbl_pr.append(layout)
            widths = (5.0, 7.0)
            for col, width in zip(t.columns, widths, strict=True):
                col.width = Cm(width)
            for row in t.rows:
                for cell, width in zip(row.cells, widths, strict=True):
                    cell.width = Cm(width)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(3)

        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading(self, text, level):
        self.doc.add_heading(text, level=level)

    def para(self, text, bold_lead=None, style=None):
        p = self.doc.add_paragraph(style=style)
        if bold_lead:
            p.add_run(bold_lead + " ").bold = True
        p.add_run(text)
        return p

    def label(self, text):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        if self.house_style:
            from docx.shared import RGBColor

            run.font.color.rgb = RGBColor.from_string(HOUSE_HEADING_RGB)
            p.paragraph_format.keep_with_next = True
        return p

    def bullet(self, text):
        return self.para(text, style="List Bullet")

    def numbered(self, text):
        return self.para(text, style="List Number")

    def caption(self, kind, text):
        from docx.shared import Pt

        if kind == "Table":
            self.n_table += 1
            n = self.n_table
        else:
            self.n_figure += 1
            n = self.n_figure
        if self.house_style:
            p = self.doc.add_paragraph(style="Caption")
            run = p.add_run(f"{kind} {n}: {text}")
            run.italic = True
            run.font.size = Pt(9)
            if kind == "Table":  # table headings sit above their table
                p.paragraph_format.keep_with_next = True
            return p
        p = self.doc.add_paragraph()
        run = p.add_run(f"{kind} {n}: {text}")
        run.italic = True
        run.font.size = Pt(9)
        return p

    def table(self, header, rows, widths_cm=None, group_rows=()):
        """header: list[str]; rows: list[list[str]]; group_rows: indices of
        rows rendered as bold shaded sub-headings spanning the table."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        def shade(cell, fill):
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tc_pr.append(shd)

        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.autofit = False
        if self.table_align == "left":
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            self._indent_to_margin(t)
        else:
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            shade(cell, HOUSE_TABLE_HEADER if self.house_style else "D9E2F3")
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            if r_i in group_rows:
                merged = cells[0].merge(cells[-1])
                merged.text = ""
                run = merged.paragraphs[0].add_run(row[0])
                run.bold = True
                run.font.size = Pt(9)
                shade(merged, "F2F2F2")
                continue
            for i, value in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(value))
                run.font.size = Pt(9)
        if self.house_style:
            self._keep_table_together(t)
        if widths_cm:
            for i, w in enumerate(widths_cm):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    if i < len(row.cells):
                        row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    @staticmethod
    def _indent_to_margin(t):
        """Indent the table by its left cell margin (108 twips in Table
        Grid), so its left border sits on the text margin: in Word's
        compatibility mode the indent is measured to the cell text."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        tbl_pr = t._tbl.tblPr
        ind = tbl_pr.find(qn("w:tblInd"))
        if ind is None:
            ind = OxmlElement("w:tblInd")
            # schema order: ... jc, tblCellSpacing, tblInd, tblBorders, shd,
            # tblLayout, tblCellMar, tblLook
            later = ("w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar",
                     "w:tblLook")  # fmt: skip
            nxt = next((e for e in tbl_pr if e.tag in {qn(x) for x in later}), None)
            if nxt is None:
                tbl_pr.append(ind)
            else:
                nxt.addprevious(ind)
        ind.set(qn("w:w"), "108")
        ind.set(qn("w:type"), "dxa")

    KEEP_TOGETHER_ROWS = 30  # tables up to this many rows are not split

    @staticmethod
    def _keep_table_together(t):
        """The header row repeats on every page, rows are not split across
        pages, and a short table (up to KEEP_TOGETHER_ROWS rows) is kept on
        one page with its heading (keep-with-next on every row but the
        last)."""
        from docx.oxml import OxmlElement

        rows = t.rows
        for i, row in enumerate(rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                tr_pr.append(OxmlElement("w:tblHeader"))
            if len(rows) <= Doc.KEEP_TOGETHER_ROWS and i < len(rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
        if len(rows) > Doc.KEEP_TOGETHER_ROWS:  # at least header + first row
            for cell in rows[0].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
        return t

    def picture(self, stream, caption, width_cm=None):
        from docx.shared import Cm

        if width_cm is None:
            width_cm = HOUSE_FIG_WIDTH_CM if self.house_style else FIG_WIDTH_CM
        self.doc.add_picture(stream, width=Cm(width_cm))
        if self.house_style:
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        self.caption("Figure", caption)
