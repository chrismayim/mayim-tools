"""Report standard v2.3: A4 page with a 16.5 cm text width, body tables on
the text margin (title-page table centred), solid grey chart gridlines."""

import pytest


def test_house_doc_page_and_table_placement(tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn

    from mayim_tools._common.docx_report import Doc

    d = Doc(house_style=True)
    d.title_page("T", "Purpose.", [("Project number", "-")])
    d.table(["A", "B"], [["1", "2"]], widths_cm=[8.0, 8.5])
    p = tmp_path / "r.docx"
    d.doc.save(str(p))
    doc = Document(str(p))
    sec = doc.sections[0]
    assert abs(sec.page_width.cm - 21.0) < 0.01
    assert (
        abs((sec.page_width - sec.left_margin - sec.right_margin) / 360000 - 16.5)
        < 0.01
    )
    jc = [t._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) for t in doc.tables]
    assert jc[0] == "center" and jc[-1] == "left"
    ind = doc.tables[-1]._tbl.tblPr.find(qn("w:tblInd"))
    assert ind.get(qn("w:w")) == "108"
    # plain Doc keeps centred tables
    assert Doc().table_align == "center"


def test_chart_gridlines_solid_grey():
    pytest.importorskip("matplotlib")
    from mayim_tools._common import report_figures as rf

    assert rf.CHART_GRID_MAJOR == "#808080" and rf.CHART_GRID_MINOR == "#D9D9D9"
    style = rf.ChartStyle()
    with style.context():
        fig = style.figure(5.0)
        ax = fig.subplots()
        ax.plot([0, 1], [0, 1])
        style.axes(ax)
        major = ax.xaxis.get_gridlines()[0]
        assert major.get_color() in ("#808080", (0.5019607843137255,) * 3)
        assert major.get_alpha() in (None, 1, 1.0)
