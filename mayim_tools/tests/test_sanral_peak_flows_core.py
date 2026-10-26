"""
Tests for mayim_tools/hydrology/sanral_peak_flows against the SANRAL
Drainage Manual Application Guide worked examples 3.1 (Moretele Spruit,
28.5 km2) and 3.2 (Tsitsa River, 4318 km2), plus checks on the
transcribed tables. Zero QGIS dependency.
"""

from __future__ import annotations

import json
import math

import pytest

from mayim_tools.hydrology.sanral_peak_flows import methods as m
from mayim_tools.hydrology.sanral_peak_flows import rainfall as rf
from mayim_tools.hydrology.sanral_peak_flows import tables as tb
from mayim_tools.hydrology.sanral_peak_flows.core import (
    CatchmentInput,
    Settings,
    field_names,
    run,
    run_catchment,
)
from mayim_tools.hydrology.sanral_peak_flows.inputs import (
    InputError,
    read_c_csv,
    read_tr102_csv,
)

# Example 3.1 catchment (AG pp 3-6 to 3-37)
A1, L1, S1, LC1, MAP1 = 28.5, 7.25, 0.02146, 4.65, 746.6
C1 = m.RunoffCoefficients(
    alpha=0.4, c1=0.337, ft_class="flat_permeable", beta=0.6, c2=0.48
)
# Pretoria (The Willows) 513524, AG Table 3.2 - Alternative 2 of example 3.1
WILLOWS = rf.TR102Station(
    label="513524 Pretoria (The Willows)",
    depths={
        1: dict(zip(tb.RETURN_PERIODS, (60, 83, 101, 121, 150, 175, 202), strict=True)),
        2: dict(
            zip(tb.RETURN_PERIODS, (75, 105, 129, 155, 192, 224, 259), strict=True)
        ),
        3: dict(
            zip(tb.RETURN_PERIODS, (83, 117, 143, 171, 211, 245, 282), strict=True)
        ),
        7: dict(
            zip(tb.RETURN_PERIODS, (110, 160, 199, 241, 303, 355, 412), strict=True)
        ),
    },
    m_mm=60.0,
    r_days=61.0,
    source="AG Table 3.2",
)


class FixedDDF:
    """Stand-in DDF giving the AG example 3.1 Alternative 3 depths at Tc."""

    def __init__(self, depths):
        self.depths = depths

    def depth_at_ari(self, duration_min, T):
        return self.depths[T], "interp"


def q(results, T):
    return next(r for r in results if r.return_period == T).q_m3s


# ------------------------------------------------------------ tables -----
def test_tr102_table_consistent_with_table_3b1():
    for basin, row in tb.TR102_BASIN_STATIONS.items():
        assert row["depths"][1][0] == tb.SDF_BASINS[basin][2]  # 2-yr 1-day == M
        for days in (1, 2, 3, 7):
            assert len(row["depths"][days]) == 7
            assert list(row["depths"][days]) == sorted(row["depths"][days])
        for i in range(7):
            col = [row["depths"][d][i] for d in (1, 2, 3, 7)]
            assert col == sorted(col)


def test_table_shapes():
    assert set(tb.SDF_BASINS) == set(range(1, 30))
    assert all(len(v[1]) == 4 for v in tb.VELD_ZONES.values())
    for region, rows in tb.RMF_RATIOS.items():
        assert set(rows) == {50, 100, 200}
        assert len({len(v) for v in rows.values()}) == 1, region
    assert tb.FT_FACTORS["flat_permeable"][20] == 0.67


# ---------------------------------------------------------- rainfall -----
def test_hershfield_reproduces_examples():
    # AG example 3.1 Alternative 2 (t = 80 min, M 60, R 61)
    assert rf.hershfield_depth(80, 20, 60, 61) == pytest.approx(85.55, abs=0.01)
    assert rf.hershfield_depth(80, 50, 60, 61) == pytest.approx(107.10, abs=0.01)
    # AG example 3.1 SDF (basin 1: M 56, R 30, Tc 1.338 h)
    assert rf.hershfield_depth(1.338 * 60, 20, 56, 30) == pytest.approx(70.84, abs=0.01)
    # AG example 3.2 SDF Table 3.27 (Tc 31.16 h, M 60, R 45)
    for T, p in ((10, 116.81), (20, 144.32), (50, 180.69), (100, 208.20)):
        assert rf.hershfield_depth(31.16 * 60, T, 60, 45) == pytest.approx(p, abs=0.01)


def test_arf_eq310_examples():
    assert rf.arf_eq310(28.5, 1.338) == pytest.approx(96.0, abs=0.1)
    assert rf.arf_eq310(4318, 31.16) == pytest.approx(80.0, abs=0.3)
    assert rf.arf_eq310(1.0, 0.25) == 100.0


def test_arf_pietersen_reproduces_part2_table2():
    # (region, A, D, T2 %, T100 %) from Pietersen et al (2024) Part 2 Table 2
    for region, area, dur, a2, a100 in (
        (1, 6120, 24.9, 64.6, 81.4),
        (2, 28893, 78.3, 74.3, 87.3),
        (4, 12854, 56.7, 73.6, 88.6),
        (5, 10260, 50.5, 74.5, 86.7),
    ):
        assert rf.arf_pietersen(region, area, dur, 2) == pytest.approx(a2, abs=0.06)
        assert rf.arf_pietersen(region, area, dur, 100) == pytest.approx(a100, abs=0.06)


def test_pietersen_falls_back_below_24h():
    arf, label, flags = rf.areal_reduction("pietersen", 28.5, 1.338, 20, 1)
    assert label == "DM Eq 3.10" and arf == pytest.approx(96.0, abs=0.1)
    assert "D < 24 h" in flags[0]


def test_tr102_interpolation_and_cap():
    st = rf.basin_station(23)
    depth, how, _ = rf.tr102_point_depth(31.16, 10, st, cap_at_1day=False)
    assert "1 - 2 day" in how
    assert depth == pytest.approx(95 + (31.16 - 24) / 24 * (121 - 95))
    # 6 - 24 h: linear from Hershfield at 6 h to the 1-day depth
    h6 = rf.hershfield_depth(360, 20, st.m_mm, st.r_days)
    depth, _, _ = rf.tr102_point_depth(15.0, 20, st, cap_at_1day=False)
    assert depth == pytest.approx(h6 + 0.5 * (111 - h6))
    # cap: Hershfield above the 1-day depth (Rational Alternative 2 only)
    capped, _, flags = rf.tr102_point_depth(5.9, 200, st, cap_at_1day=True)
    raw, _, _ = rf.tr102_point_depth(5.9, 200, st, cap_at_1day=False)
    assert raw > 175 and capped == 175 and flags


# --------------------------------------------------- example 3.1 ---------
def test_runoff_coefficient_example_3_1():
    assert C1.c_t(20) == pytest.approx(0.3783, abs=0.0001)
    assert C1.c_t(50) == pytest.approx(0.3999, abs=0.0001)


def test_rational_alternative_2_example_3_1():
    res = m.rational(
        2,
        A1,
        1.338,
        C1,
        (20, 50),
        tc_floor=True,
        apply_arf=True,
        arf_method="eq310",
        arf_region=None,
        station=WILLOWS,
    )
    i20 = res[0].derivation["I_mmh"]
    i50 = res[1].derivation["I_mmh"]
    assert i20 == pytest.approx(61.4, abs=0.15)
    assert i50 == pytest.approx(76.9, abs=0.15)
    assert q(res, 20) == pytest.approx(184, abs=1.0)
    assert q(res, 50) == pytest.approx(243, abs=1.0)
    assert any("Table 3.2" in f for f in res[0].flags)  # 28.5 > 15 km2


def test_rational_alternative_3_example_3_1_without_arf():
    ddf = FixedDDF({20: 57.0, 50: 70.0})
    res = m.rational(
        3,
        A1,
        1.338,
        C1,
        (20, 50),
        tc_floor=True,
        apply_arf=False,
        arf_method="eq310",
        arf_region=None,
        ddf=ddf,
    )
    assert q(res, 20) == pytest.approx(128, abs=0.6)
    assert q(res, 50) == pytest.approx(166, abs=0.6)


def test_tc_floor_switch():
    ddf = FixedDDF({20: 20.0})
    kw = dict(apply_arf=False, arf_method="eq310", arf_region=None, ddf=ddf)
    on = m.rational(3, 1.0, 0.1, C1, (20,), tc_floor=True, **kw)[0]
    off = m.rational(3, 1.0, 0.1, C1, (20,), tc_floor=False, **kw)[0]
    assert on.derivation["Tc_used_h"] == 0.25 and off.derivation["Tc_used_h"] == 0.1
    assert on.q_m3s == pytest.approx(off.q_m3s * 0.1 / 0.25)


def test_sdf_example_3_1():
    res = m.sdf(A1, L1, S1, 1, rf.basin_station(1), (20, 50))
    assert res[0].derivation["Tc_h"] == pytest.approx(1.338, abs=0.001)
    assert res[0].derivation["C_T"] == pytest.approx(0.3112, abs=0.0001)
    assert res[1].derivation["C_T"] == pytest.approx(0.3639, abs=0.0001)
    assert res[0].derivation["I_mmh"] == pytest.approx(50.83, abs=0.05)
    assert q(res, 20) == pytest.approx(125, abs=0.6)
    assert q(res, 50) == pytest.approx(183, abs=0.6)


def test_midgley_pitman_example_3_1():
    res = m.midgley_pitman(A1, MAP1, L1, LC1, S1, "8")
    assert res[0].derivation["C_param"] == pytest.approx(0.1238, abs=0.0001)
    assert q(res, 20) == pytest.approx(78.85, abs=0.01)
    assert q(res, 50) == pytest.approx(109.28, abs=0.01)


# --------------------------------------------------- example 3.2 ---------
A2, L2, S2, LC2, MAP2 = 4318.0, 179.5, 0.0037, 85.0, 860.0


def test_sdf_example_3_2_coefficients_and_hershfield_path():
    """AG Table 3.27 used Hershfield at Tc = 31.2 h; the chain C_T, I, Q is
    checked with those depths (ARF 80 % as tabulated)."""
    tc = m.sdf_tc_hours(L2, S2)
    assert tc == pytest.approx(31.16, abs=0.01)
    table = {10: (0.485, 1739), 20: (0.593, 2628), 50: (0.716, 3974), 100: (0.80, 5117)}
    for T, (ct, qt) in table.items():
        c_t = 0.10 + tb.SDF_YT[T] / 2.33 * (0.80 - 0.10)
        assert c_t == pytest.approx(ct, abs=0.001)
        p = rf.hershfield_depth(tc * 60, T, 60, 45)
        assert c_t * p * 0.80 / tc * A2 / 3.6 == pytest.approx(qt, rel=0.003)


def test_sdf_example_3_2_follows_manual_step_6():
    """Chris's decision (option a): above 24 h the Manual's TR102
    interpolation is used, not the guide's Hershfield."""
    res = m.sdf(A2, L2, S2, 23, rf.basin_station(23), (10, 100))
    assert "TR102 1 - 2 day" in res[0].derivation["rain_source"]
    assert res[0].derivation["P_mm"] == pytest.approx(102.76, abs=0.01)
    assert q(res, 10) < 1739


def test_midgley_pitman_example_3_2():
    res = m.midgley_pitman(A2, MAP2, L2, LC2, S2, "2A")  # as used in the AG
    assert res[0].derivation["C_param"] == pytest.approx(0.01721, abs=0.00001)
    for T, qt in ((10, 1812), (20, 2270), (50, 2969), (100, 3493)):
        assert q(res, T) == pytest.approx(qt, abs=1.5)


def test_rmf_and_ratios_example_3_2():
    r = m.rmf(A2, "K6", None)
    assert r.q_m3s == pytest.approx(8059, abs=1)
    ratios = m.rmf_ratios(A2, "K6", r.q_m3s, (50, 100, 200))
    assert q(ratios, 50) == pytest.approx(3951, abs=2)
    assert q(ratios, 100) == pytest.approx(4796, abs=2)


def test_rmf_zones_and_custom_k():
    assert m.rmf(50, "K5", None).derivation["equation"].endswith("(transition zone)")
    flood = m.rmf(1000, "K5", None)
    assert flood.q_m3s == pytest.approx(100 * 1000**0.5)
    custom = m.rmf(1000, None, 5.0)
    assert custom.q_m3s == pytest.approx(1e6 * (1000 / 1e8) ** 0.5)


def test_rmf_ratio_limits():
    ratio, flags = m.rmf_ratio_value("K8", 50, 20000)
    assert ratio is None and flags
    ratio, _ = m.rmf_ratio_value("K8", 50, 5)
    assert ratio == 0.537
    assert (
        m.rmf_ratio_value("K1", 100, 1000)[0] == m.rmf_ratio_value("K2", 100, 1000)[0]
    )


# ------------------------------------------------------------- CAPA ------
def test_capa_digitised_lines():
    # the M lines reproduce their own digitised points
    for mm, c in tb.CAPA_INTERCEPTS:
        assert m._capa_intercept(mm) == pytest.approx(c)
    # read-back check against the diagram: M = 10 000 at A = 100 -> ~117 m3/s
    q50 = 10 ** (tb.CAPA_SLOPE * 2 + m._capa_intercept(10000))
    assert 105 < q50 < 125


def test_capa_kp_and_growth():
    res = m.capa(500.0, 800.0, 0.05, 40.0, (2, 5, 50, 100))
    q50 = res[0].q_m3s
    assert res[0].return_period == 2
    assert q(res, 50) == pytest.approx(5.46 * q50)
    assert all("digitised" in f for r in res for f in r.flags[-1:])


def test_capa_outside_diagram():
    res = m.capa(500.0, 50.0, 0.001, 400.0, (2, 50))
    assert all(r.q_m3s is None for r in res)


# ------------------------------------------------------ orchestration ----
def _catchment(**kw):
    base = dict(
        catchment_id="C1",
        area_km2=A1,
        lfp_km=L1,
        s1085=S1,
        lca_km=LC1,
        s_mean=0.08,
        tc_min=1.338 * 60,
        map_mm=MAP1,
        sdf_basin=1,
        veld_zone="8",
        kovacs_region="K5",
    )
    base.update(kw)
    return CatchmentInput(**base)


def test_run_catchment_all_methods_and_fields():
    s = Settings()
    cr = run_catchment(
        _catchment(),
        s,
        ddf=FixedDDF({T: 50.0 + T for T in tb.RETURN_PERIODS}),
        c_values={"C1": C1},
    )
    got = set(cr.by_field())
    assert got == set(field_names(s))
    assert "Q_RMF" in got and "Q_MP_10" in got and "Q_RMFR_200" in got
    assert "Q_CAPA_2" in got and "Q_MP_2" not in got


def test_run_skips_with_notes():
    cr = run_catchment(_catchment(sdf_basin=None, kovacs_region=None), Settings())
    text = cr.notes_text()
    assert "rational_alt2: not calculated" in text
    assert "sdf: not calculated - no SDF basin" in text
    assert "rmf: not calculated" in text


def test_incremental_and_ineffective_area_notes():
    cr = run_catchment(
        _catchment(mode="incremental", ineffective_km2=1.5), Settings(methods=("sdf",))
    )
    assert any("incremental" in n for n in cr.notes)
    assert cr.catchment.effective_area == pytest.approx(27.0)


def test_run_warns_on_unmatched_c_rows():
    res = run([_catchment()], Settings(methods=("sdf",)), c_values={"C1": C1, "X": C1})
    assert "X" in res.warnings[0]


# ------------------------------------------------------------- inputs ----
def test_read_c_csv(tmp_path):
    p = tmp_path / "c.csv"
    p.write_text(
        "catchment_id;alpha;C1;Ft_class;beta;C2;gamma;C3\n"
        "C1;0,4;0,337;flat;0,6;0,48;;\n"
        "C2;1;0.5;steep_impermeable;;;;\n",
        encoding="utf-8",
    )
    vals = read_c_csv(p)
    assert vals["C1"].c_t(20) == pytest.approx(0.3783, abs=0.0001)
    assert vals["C2"].ft_class == "steep_impermeable" and vals["C2"].beta == 0.0


def test_read_c_csv_errors(tmp_path):
    p = tmp_path / "c.csv"
    p.write_text("catchment_id,alpha,C1,Ft_class\nC1,1,0.3,sandy\n", encoding="utf-8")
    with pytest.raises(InputError):
        read_c_csv(p)


def test_read_tr102_csv(tmp_path):
    p = tmp_path / "t.csv"
    rows = [
        "Station,513524 Pretoria (The Willows)",
        "R,61",
        "Duration (days),2,5,10,20,50,100,200",
    ]
    for d in (1, 2, 3, 7):
        rows.append(
            ",".join([str(d)] + [str(WILLOWS.depths[d][T]) for T in tb.RETURN_PERIODS])
        )
    p.write_text("\n".join(rows) + "\n", encoding="utf-8")
    st = read_tr102_csv(p, None)
    assert st.m_mm == 60 and st.r_days == 61 and st.depth(7, 200) == 412


# ------------------------------------------------------------ exports ----
def _full_run():
    s = Settings()
    return run(
        [
            _catchment(),
            _catchment(
                catchment_id="C2",
                area_km2=4318.0,
                lfp_km=179.5,
                s1085=0.0037,
                lca_km=85.0,
                tc_min=31.16 * 60,
                map_mm=860.0,
                sdf_basin=23,
                veld_zone="5",
                kovacs_region="K6",
            ),
        ],
        s,
        ddf=FixedDDF({T: 50.0 + T for T in tb.RETURN_PERIODS}),
        c_values={"C1": C1},
    )


def test_csv_export(tmp_path):
    from mayim_tools.hydrology.sanral_peak_flows.export import write_csv

    res = _full_run()
    p = tmp_path / "out.csv"
    write_csv(p, res, {"tool": "test", "version": "0"})
    text = p.read_text(encoding="utf-8-sig")
    for n in range(1, 10):
        assert f"# {n} " in text
    assert "Q_SDF_100" in text and "8059.42" in text
    assert (
        "," not in text.split("# 4")[1].splitlines()[2].split(",")[1]
    )  # no 1,000 separators


def test_results_json(tmp_path):
    from mayim_tools._common.results_json import SCHEMA, read_results
    from mayim_tools.hydrology.sanral_peak_flows.export import write_results_json

    res = _full_run()
    p = tmp_path / "results.json"
    write_results_json(p, res, {"id": "sanral_peak_flows"}, {"files": {}})
    data = read_results(p)
    assert data["schema"] == SCHEMA and len(data["items"]) == 2
    big = data["items"][1]["results"]
    rmf = next(r for r in big if r["method"] == "rmf")
    assert rmf["return_period_years"] is None and rmf["aep_percent"] is None
    sdf100 = next(
        r for r in big if r["method"] == "sdf" and r["return_period_years"] == 100
    )
    assert sdf100["aep_percent"] == 1.0 and sdf100["unit"] == "m3/s"
    json.dumps(data)  # serialisable


def test_word_report(tmp_path):
    pytest.importorskip("docx")
    pytest.importorskip("matplotlib")
    import re

    from docx import Document

    from mayim_tools.hydrology.sanral_peak_flows.report import num, write_report

    assert num(1234.5, 1) == "1234.5"  # decimal point, no thousands separator
    p = tmp_path / "annex.docx"
    write_report(
        p,
        _full_run(),
        {"tool": "SANRAL Deterministic Peak Flows", "version": "0.1.0"},
        context={
            "project_number": "MC2025_014",
            "project_name": "Riverside Estate",
            "outputs": [("Peak flow calculations (CSV)", "D:/p/x.csv")],
        },
    )
    assert p.stat().st_size > 10000
    doc = Document(str(p))
    heads = [x.text for x in doc.paragraphs if x.style.name == "Heading 1"]
    assert heads == [
        "Summary",
        "Tool and method",
        "Inputs and settings",
        "Parameters and calculations",
        "Results",
        "Outputs",
        "Notes and limitations",
        "References",
    ]
    text = "\n".join(x.text for x in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "MC2025_014" in cells and "Table 1:" in text
    assert not re.search(r"\d,\d{3}", cells)  # no thousands separators
    assert not re.search(r"\d,\d", cells)  # no decimal commas
