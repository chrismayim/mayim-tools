"""Time of Concentration core: methods hand-checked, the path split, slopes,
land-cover suggestions and the selected Tc."""

import math

import numpy as np
import pytest

from mayim_tools.hydrology.time_of_concentration import tables as tb
from mayim_tools.hydrology.time_of_concentration.core import (
    CatchmentTc,
    Params,
    Settings,
    calculate,
    fill_parameters,
    profile_from_line,
    slope_10_85,
    suggest,
    tc_bransby_williams_minutes,
    tc_kerby_minutes,
    tc_kirpich_minutes,
    tc_shallow_minutes,
    tc_usbr_minutes,
)
from mayim_tools.hydrology.time_of_concentration.export import write_csv


def test_empirical_formulas_hand_checked():
    # L = 1 km, S = 0.01 m/m
    assert tc_kirpich_minutes(1000.0, 0.01) == pytest.approx(23.45, abs=0.05)
    assert tc_usbr_minutes(1.0, 0.01) == pytest.approx(
        60 * (0.87 / 10.0) ** 0.385, rel=1e-9
    )
    assert tc_usbr_minutes(1.0, 0.01) == pytest.approx(23.44, abs=0.05)
    assert tc_bransby_williams_minutes(2.0, 4.0, 0.005) == pytest.approx(
        58 * 2.0 / (4.0**0.1 * 5.0**0.2)
    )
    assert tc_kirpich_minutes(1000.0, 0.0) is None
    assert tc_usbr_minutes(1.0, None) is None
    # USBR (Eq 3.12) is Kirpich in other units (same coefficient to 0.1 %)
    ratio = tc_usbr_minutes(2.5, 0.02) / tc_kirpich_minutes(2500.0, 0.02)
    assert ratio == pytest.approx(1.0, abs=0.001)


def test_kerby_matches_dm_eq_3_11_and_txdot_si():
    # DM: Tc = 0.604 (r L / S^0.5)^0.467 h, L km
    t = tc_kerby_minutes(100.0, 0.4, 0.05)
    assert t == pytest.approx(60 * 0.604 * (0.4 * 0.1 / math.sqrt(0.05)) ** 0.467)
    # TxDOT SI form: t = 1.44 (L N)^0.467 S^-0.235 min, L m
    assert t == pytest.approx(1.44 * (100 * 0.4) ** 0.467 * 0.05**-0.235, rel=0.01)
    assert t == pytest.approx(16.2, abs=0.2)


def test_shallow_flow_neh_coefficients():
    # NEH 630 Table 15-3 short-grass pasture: V = 6.9620 S^0.5 ft/s
    k = tb.SC_TYPES["pasture"][1]
    assert k == pytest.approx(6.962 * 0.3048)
    # 500 m at 4 %: V = 0.424 m/s, 19.6 min
    assert tc_shallow_minutes(500.0, k, 0.04) == pytest.approx(500 / (k * 0.2) / 60)
    assert tc_shallow_minutes(500.0, k, 0.04) == pytest.approx(19.6, abs=0.1)


def test_profile_from_line_reverses_to_outlet_distance():
    xy = [(0, 0), (30, 0), (60, 0), (90, 0)]  # divide -> outlet
    x, z = profile_from_line(xy, [13.0, np.nan, 11.0, 10.0])
    assert list(x) == [0.0, 30.0, 90.0]
    assert list(z) == [10.0, 11.0, 13.0]


def _linear_catchment(**params):
    # 2 km path: channel 0-1500 m at 1 %, then 3 % to the divide
    x = np.linspace(0.0, 2000.0, 201)
    z = np.where(x <= 1500, 0.01 * x, 15.0 + 0.03 * (x - 1500))
    return CatchmentTc(
        "A", 1.0, 2.0, s_avg=float(z[-1] / 2000), s_1085=0.015, s_ea=0.012,
        ch_head_m=1500.0, x=x, z=z, params=Params(**params),
    )  # fmt: skip


def test_combined_split_slopes_and_sum():
    c = _linear_catchment(ov_r=0.4, sc_type="pasture")
    fill_parameters(c, Settings())
    assert c.params.ch_start_m == 1500.0 and c.params.ov_len_m == 100.0
    calculate(c, Settings())
    v = c.values
    assert (v["l_ov_m"], v["l_sc_m"], v["l_ch_m"]) == (100.0, 400.0, 1500.0)
    assert v["s_ov"] == pytest.approx(0.03) and v["s_sc"] == pytest.approx(0.03)
    assert v["s_ch"] == pytest.approx(0.01)
    assert v["tc_ov_mn"] == pytest.approx(tc_kerby_minutes(100, 0.4, 0.03))
    assert v["tc_sc_mn"] == pytest.approx(
        tc_shallow_minutes(400, tb.SC_TYPES["pasture"][1], 0.03)
    )
    assert v["tc_ch_mn"] == pytest.approx(tc_usbr_minutes(1.5, 0.01))
    assert v["tc_comb_mn"] == pytest.approx(
        v["tc_ov_mn"] + v["tc_sc_mn"] + v["tc_ch_mn"]
    )
    assert v["tc_sel_mn"] == v["tc_comb_mn"] and v["tc_sel_by"] == "combined"
    assert not c.missing


def test_missing_parameters_leave_combined_empty():
    c = _linear_catchment()
    s = Settings()
    fill_parameters(c, s)
    calculate(c, s)
    assert set(c.missing) == {"ov_r", "sc_type"}
    assert c.values["tc_comb_mn"] is None and c.values["tc_sel_mn"] is None
    assert c.values["tc_ch_mn"] is not None  # the channel needs no parameter
    assert any("selected method combined has no value" in f for f in c.flags)


def test_overland_cut_at_channel_and_kerby_limit_and_slope_floor():
    c = _linear_catchment(ov_len_m=800.0, ov_r=0.4, sc_type="pasture")
    calculate(c, Settings())
    assert c.values["l_ov_m"] == 500.0 and c.values["l_sc_m"] == 0.0
    assert any("exceeds the 366 m" in f for f in c.flags)
    # flat channel: slope raised to the floor
    x = np.linspace(0, 1000, 101)
    flat = CatchmentTc("F", 1.0, 1.0, ch_head_m=900.0, x=x, z=np.zeros_like(x),
                       params=Params(ov_len_m=100, ov_r=0.4))  # fmt: skip
    fill_parameters(flat, Settings())
    calculate(flat, Settings())
    assert flat.values["s_ch"] == tb.S_MIN_CHANNEL
    assert any("channel slope" in f for f in flat.flags)


def test_no_channel_and_user_method():
    c = _linear_catchment(ov_r=0.4, sc_type="bare", tc_method="user",
                          tc_user_mn=42.0)  # fmt: skip
    c.ch_head_m = None
    fill_parameters(c, Settings())
    calculate(c, Settings())
    assert c.values["l_ch_m"] == 0.0 and c.values["tc_ch_mn"] is None
    assert c.values["l_sc_m"] == pytest.approx(1900.0)
    assert c.values["tc_sel_mn"] == 42.0
    c2 = _linear_catchment(tc_method="usbr")
    calculate(c2, Settings(default_ov_len_m=0))
    assert c2.values["tc_sel_mn"] == c2.values["tc_usbr_mn"]
    assert "ov_len_m" in c2.missing


def test_land_cover_suggestions_and_fill():
    assert suggest([13, 13, 40, 19], "sanlc")[:2] == (0.4, "pasture")
    r, sc, code, share = suggest([19, 19, 19, 66], "sanlc")  # water skipped
    assert (r, sc, code, share) == (0.02, "paved", 66, 25.0)
    assert suggest([80, 80], "worldcover") == (None, None, None, None)
    c = _linear_catchment()
    c.lc_overland, c.lc_shallow = [1, 1, 13], [40, 40]
    notes = fill_parameters(c, Settings(fill_from_land_cover=True))
    assert c.params.ov_r == 0.8 and c.params.sc_type == "cultivated"
    assert any("land-cover" in n for n in notes)
    calculate(c, Settings())
    assert c.values["sug_r"] == 0.8 and c.values["lc_ov"] == 1
    # without the fill option the suggestion is reported but not used
    c = _linear_catchment()
    c.lc_overland = [13]
    fill_parameters(c, Settings())
    calculate(c, Settings())
    assert c.params.ov_r is None and c.values["sug_r"] == 0.4


def test_slope_10_85_on_reach():
    x = np.linspace(0, 1000, 11)
    z = 0.02 * x
    assert slope_10_85(x, z, 0.0, 500.0) == pytest.approx(0.02)


def test_csv(tmp_path):
    c = _linear_catchment(ov_r=0.4, sc_type="pasture")
    fill_parameters(c, Settings())
    calculate(c, Settings())
    p = tmp_path / "tc.csv"
    write_csv(p, [c], [("Tool", "Time of Concentration 0.1.0")], "sanlc")
    text = p.read_text(encoding="utf-8-sig")
    assert text.startswith("# Column key") and "[catchments]" in text
    assert "tc_comb_mn" in text and "[land-cover suggestions]" in text


def test_word_report(tmp_path):
    pytest.importorskip("docx")
    pytest.importorskip("matplotlib")
    import re

    from docx import Document

    from mayim_tools.hydrology.time_of_concentration.report import write_report

    done = _linear_catchment(ov_r=0.4, sc_type="pasture")
    todo = _linear_catchment()
    todo.catchment_id = "B"
    s = Settings()
    for c in (done, todo):
        fill_parameters(c, s)
        calculate(c, s)
    p = tmp_path / "tc.docx"
    write_report(
        p,
        [done, todo],
        {"tool": "Time of Concentration", "version": "0.1.0",
         "rows": [("Catchment layer", "x.shp"),
                  ("Default overland length (m)", 100.0)]},
        context={"project_number": "MC2025_014", "settings": s,
                 "outputs": [("Calculations (CSV)", "D:/p/MC2025_014_Tc.csv")]},
    )  # fmt: skip
    doc = Document(str(p))
    heads = [x.text for x in doc.paragraphs if x.style.name == "Heading 1"]
    assert heads == [
        "Summary",
        "Tool and method",
        "Inputs and settings",
        "Parameters and calculations",
        "Results",
        "Outputs",
        "Notes and limitations",
        "References",
    ]
    text = "\n".join(x.text for x in doc.paragraphs)
    assert "B: ov_r, sc_type" in text and "Design storm duration" in text
    assert len(doc.inline_shapes) == 2
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert not re.search(r"\d,\d", cells)
    # body tables on the text margin (title-page table stays centred)
    from docx.oxml.ns import qn

    body = doc.tables[1:]
    assert all(t._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) == "left" for t in body)
    assert all(t._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")) == "108" for t in body)
    # A4 with a 16.5 cm text width
    sec = doc.sections[0]
    text_w = sec.page_width - sec.left_margin - sec.right_margin
    assert abs(text_w / 360000 - 16.5) < 0.01 and abs(sec.page_width.cm - 21.0) < 0.01
