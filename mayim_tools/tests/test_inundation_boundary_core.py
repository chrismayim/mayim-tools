"""
Tests for mayim_tools/hydrology/inundation_boundary/core.py and export.py.
Zero QGIS dependency - run under plain pytest.

Areas and volumes are checked against analytic shapes (cone, box pits);
the boundary rings against the analytic contour radius; polygon nesting
with dry islands; the two inundation modes; bounding-polygon clipping.
"""

from __future__ import annotations

import csv
import math

import numpy as np
import pytest

from mayim_tools.hydrology.inundation_boundary.core import (
    InundationError,
    fill_field,
    level_boundaries,
    nest_rings,
    run_inundation,
    seed_cells,
    water_levels,
)
from mayim_tools.hydrology.inundation_boundary.export import (
    SUMMARY_FIELDS,
    write_summary_csv,
)

CELL = 1.0
GT = (0.0, CELL, 0.0, 100.0, 0.0, -CELL)  # origin top-left at (0, 100)


def _cone(n=101, z0=100.0, slope=0.1):
    """Cone with its apex (lowest point) at the grid centre."""
    r, c = np.mgrid[0:n, 0:n]
    centre = (n - 1) / 2
    d = np.hypot(r - centre, c - centre) * CELL
    return z0 + slope * d, centre


def _centre_xy(centre):
    return GT[0] + (centre + 0.5) * CELL, GT[3] - (centre + 0.5) * CELL


def _ring_area(ring):
    x, y = np.asarray(ring)[:, 0], np.asarray(ring)[:, 1]
    return 0.5 * abs(float(np.sum(x[:-1] * y[1:] - x[1:] * y[:-1])))


def _two_pits(n=60):
    """Flat ground at 10 m with two separate 2 m deep box pits."""
    z = np.full((n, n), 10.0)
    z[10:20, 10:20] = 8.0  # 100 cells
    z[30:50, 30:45] = 8.0  # 300 cells
    return z


# ---------------------------------------------------------------------------
# Levels
# ---------------------------------------------------------------------------


def test_single_level():
    assert water_levels(101.25) == [101.25]
    assert water_levels(101.25, 101.25, 0.5) == [101.25]


def test_range_includes_end():
    assert water_levels(100.0, 101.0, 0.25) == [100.0, 100.25, 100.5, 100.75, 101.0]
    assert water_levels(100.0, 101.1, 0.5) == [100.0, 100.5, 101.0, 101.1]


def test_range_reversed_and_errors():
    assert water_levels(101.0, 100.0, 0.5) == [100.0, 100.5, 101.0]
    with pytest.raises(InundationError):
        water_levels(100.0, 101.0, 0.0)
    with pytest.raises(InundationError):
        water_levels(0.0, 10000.0, 0.001)


# ---------------------------------------------------------------------------
# Fill field and seeds
# ---------------------------------------------------------------------------


def test_all_mode_floods_both_pits():
    z = _two_pits()
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [9.0])
    lv = res.levels[0]
    assert lv.wet_cells == 400
    assert lv.area_m2 == pytest.approx(400.0)
    assert lv.volume_m3 == pytest.approx(400.0)
    assert lv.max_depth == pytest.approx(1.0)
    assert len(lv.polygons) == 2


def test_connected_mode_floods_only_seeded_pit():
    z = _two_pits()
    x, y = GT[0] + 15.5, GT[3] - 15.5  # inside the small pit
    res = run_inundation(
        z, -9999.0, GT, (0, 0, *z.shape), [9.0], mode="connected", seeds=[(x, y)]
    )
    lv = res.levels[0]
    assert res.n_seeds == 1
    assert lv.wet_cells == 100
    assert len(lv.polygons) == 1
    assert not res.warnings


def test_connected_mode_without_seed_uses_lowest_cell():
    z = _two_pits()
    z[40, 40] = 7.0  # the big pit holds the lowest cell
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [9.0], mode="connected")
    assert res.levels[0].wet_cells == 300
    assert any("lowest cell" in w for w in res.warnings)


def test_connected_through_saddle_once_level_rises():
    z = _two_pits()
    z[20:30, 15] = 9.5  # channel joining the pits, sill at 9.5
    z[30, 15:31] = 9.5
    seed = (GT[0] + 15.5, GT[3] - 15.5)
    low = run_inundation(
        z, -9999.0, GT, (0, 0, *z.shape), [9.0, 9.75], mode="connected", seeds=[seed]
    )
    assert low.levels[0].wet_cells == 100
    assert low.levels[1].wet_cells > 400  # both pits plus the channel


def test_seed_outside_domain_is_reported():
    z = _two_pits()
    cells, outside = seed_cells([(500.0, 500.0), (15.5, 84.5)], GT, z < 100)
    assert cells == [(15, 15)]
    assert outside == 1


def test_fill_field_modes():
    z = _two_pits()
    dom = np.ones_like(z, dtype=bool)
    assert np.array_equal(fill_field(z, dom, "all"), z)
    f = fill_field(z, dom, "connected", [(15, 15)])
    assert f[15, 15] == pytest.approx(8.0)
    assert f[40, 40] == pytest.approx(10.0)  # reached over the 10 m flat
    with pytest.raises(InundationError):
        fill_field(z, dom, "bogus")


# ---------------------------------------------------------------------------
# Boundaries
# ---------------------------------------------------------------------------


def test_cone_contour_matches_analytic_circle():
    z, centre = _cone()
    level = 102.0  # radius 20 m
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [level], smoothing=3)
    lv = res.levels[0]
    assert len(lv.lines) == 1 and len(lv.polygons) == 1
    ring = np.asarray(lv.lines[0])
    assert np.allclose(ring[0], ring[-1])
    cx, cy = _centre_xy(centre)
    radii = np.hypot(ring[:, 0] - cx, ring[:, 1] - cy)
    assert radii.mean() == pytest.approx(20.0, abs=0.3)
    assert _ring_area(ring) == pytest.approx(math.pi * 400, rel=0.02)
    # cell-based area and the cone volume pi r^2 h / 3
    assert lv.area_m2 == pytest.approx(math.pi * 400, rel=0.03)
    assert lv.volume_m3 == pytest.approx(math.pi * 400 * 2 / 3, rel=0.03)
    assert lv.max_depth == pytest.approx(2.0, abs=0.01)


def test_island_becomes_hole():
    z = np.full((60, 60), 10.0)
    z[10:50, 10:50] = 5.0
    z[25:35, 25:35] = 12.0  # island
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [8.0])
    lv = res.levels[0]
    assert len(lv.lines) == 2
    assert len(lv.polygons) == 1 and len(lv.polygons[0]) == 2
    outer, hole = lv.polygons[0]
    assert _ring_area(outer) > _ring_area(hole)
    assert lv.wet_cells == 1600 - 100


def test_nest_rings_three_levels():
    def square(h):
        return np.array([(-h, -h), (h, -h), (h, h), (-h, h), (-h, -h)], float)

    polys = nest_rings([square(1), square(10), square(5)])
    # 10 = water with hole 5; 1 = water again (lake on the island)
    assert len(polys) == 2
    assert len(polys[0]) == 2 and _ring_area(polys[0][1]) == pytest.approx(100)
    assert len(polys[1]) == 1 and _ring_area(polys[1][0]) == pytest.approx(4)


def test_small_rings_dropped():
    z = np.full((30, 30), 10.0)
    z[15, 15] = 5.0  # single-cell pit
    rings, polys = level_boundaries(z, 9.0, GT, 0, 3.0)
    assert rings == [] and polys == []
    rings, _ = level_boundaries(z, 9.0, GT, 0, 0.0)
    assert len(rings) == 1


def test_no_smoothing_ring_through_midpoints():
    z = np.full((20, 20), 10.0)
    z[5:10, 5:10] = 0.0
    rings, _ = level_boundaries(z, 5.0, GT, smoothing=0, min_cells=0)
    ring = np.asarray(rings[0])
    # edge exactly half way between the pit and flat cell centres
    assert ring[:, 0].min() == pytest.approx(GT[0] + 5.0)
    assert ring[:, 0].max() == pytest.approx(GT[0] + 10.0)


# ---------------------------------------------------------------------------
# Bounding polygon
# ---------------------------------------------------------------------------


def test_bounding_polygon_limits_and_closes_rings():
    z, centre = _cone()
    cx, cy = _centre_xy(centre)
    # half of the circle: the polygon covers x <= cx only
    bound = [
        [
            [
                (cx - 60, cy - 60),
                (cx, cy - 60),
                (cx, cy + 60),
                (cx - 60, cy + 60),
                (cx - 60, cy - 60),
            ]
        ]
    ]
    full = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [102.0]).levels[0]
    half = run_inundation(
        z, -9999.0, GT, (0, 0, *z.shape), [102.0], bounding=bound
    ).levels[0]
    assert half.area_m2 == pytest.approx(full.area_m2 / 2, rel=0.03)
    assert half.volume_m3 == pytest.approx(full.volume_m3 / 2, rel=0.03)
    # the ring is closed and extends at most a cell beyond the polygon edge
    ring = np.asarray(half.lines[0])
    assert np.allclose(ring[0], ring[-1])
    assert ring[:, 0].max() <= cx + CELL


def test_bounding_off_dem_raises():
    z = _two_pits()
    bound = [[[(500, 500), (510, 500), (510, 510), (500, 510), (500, 500)]]]
    with pytest.raises(InundationError):
        run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [9.0], bounding=bound)


def test_nodata_is_dry():
    z = _two_pits()
    z[12:18, 12:18] = -9999.0
    lv = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [9.0]).levels[0]
    assert lv.wet_cells == 400 - 36
    assert len(lv.polygons) == 2
    assert len(lv.polygons[0]) + len(lv.polygons[1]) == 3  # NoData hole


# ---------------------------------------------------------------------------
# Range, warnings, depth raster, CSV
# ---------------------------------------------------------------------------


def test_range_monotonic_and_depth_raster():
    z, _ = _cone()
    levels = water_levels(100.5, 103.0, 0.5)
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), levels)
    areas = [lv.area_m2 for lv in res.levels]
    vols = [lv.volume_m3 for lv in res.levels]
    assert areas == sorted(areas) and vols == sorted(vols)
    assert res.depth_level == pytest.approx(103.0)
    assert np.nanmax(res.depth) == pytest.approx(3.0, abs=0.01)
    assert np.nansum(res.depth) * CELL * CELL == pytest.approx(vols[-1], rel=1e-6)


def test_warnings_dry_and_overtopped():
    z = _two_pits()
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [7.0, 11.0])
    assert res.levels[0].wet_cells == 0
    assert any("nothing is inundated" in w for w in res.warnings)
    assert any("at or above the highest ground" in w for w in res.warnings)


def test_summary_csv(tmp_path):
    z = _two_pits()
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [8.5, 9.0])
    path = tmp_path / "summary.csv"
    write_summary_csv(str(path), res, {9.0: (401.2, 80.4, 2)})
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == SUMMARY_FIELDS
    assert rows[0]["level_m"] == "8.500" and rows[0]["n_polygons"] == ""
    assert rows[1]["volume_m3"] == "400.0"
    assert rows[1]["mean_depth_m"] == "1.000"
    assert rows[1]["n_polygons"] == "2"
    assert "," not in rows[1]["area_m2"]  # no thousands separators


def test_per_level_depth_and_water_surface():
    z = _two_pits()
    z[12, 12] = 7.5
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [8.5, 9.0])
    d85, w85 = res.depth_at(8.5), res.water_surface_at(8.5)
    assert np.nanmax(d85) == pytest.approx(1.0)
    assert np.count_nonzero(np.isfinite(d85)) == 400
    assert np.nanmin(w85) == np.nanmax(w85) == pytest.approx(8.5)
    # water surface only where depth is defined
    assert np.array_equal(np.isfinite(w85), np.isfinite(d85))
    assert np.array_equal(res.depth_at(9.0), res.depth, equal_nan=True)


def test_seed_above_level_is_explained():
    z = _two_pits()
    bank = (GT[0] + 5.5, GT[3] - 5.5)  # on the 10 m flat, not in a pit
    res = run_inundation(
        z, -9999.0, GT, (0, 0, *z.shape), [9.0], mode="connected", seeds=[bank]
    )
    assert res.levels[0].wet_cells == 0
    assert res.seed_ground == [10.0]
    assert any("ground at the seed point" in w for w in res.warnings)


# ---------------------------------------------------------------------------
# Word report
# ---------------------------------------------------------------------------


def test_depth_classes_sum_to_area():
    from mayim_tools.hydrology.inundation_boundary.report import depth_classes

    z, _ = _cone()
    res = run_inundation(z, -9999.0, GT, (0, 0, *z.shape), [103.0])
    classes = depth_classes(res.depth_at(103.0), 1.0, res.weights)
    assert sum(c["area_m2"] for c in classes) == pytest.approx(res.levels[0].area_m2)
    assert sum(c["pct"] for c in classes) == pytest.approx(100.0)
    assert classes[-1]["to_m"] >= 3.0 - 1e-9


def test_report_docx(tmp_path):
    pytest.importorskip("docx")
    pytest.importorskip("matplotlib")
    import docx

    from mayim_tools.hydrology.inundation_boundary.report import write_report_docx

    z, centre = _cone()
    cx, cy = _centre_xy(centre)
    bound = [[[(cx - 40, cy - 40), (cx + 40, cy - 40), (cx + 40, cy + 40),
               (cx - 40, cy + 40), (cx - 40, cy - 40)]]]  # fmt: skip
    levels = water_levels(99.5, 103.0, 0.5)
    res = run_inundation(
        z, -9999.0, GT, (0, 0, *z.shape), levels, bounding=bound,
        mode="connected", seeds=[(cx, cy)],
    )  # fmt: skip
    geometry = {lv.level: (lv.area_m2, 100.0, len(lv.polygons)) for lv in res.levels}
    path = tmp_path / "report.docx"
    write_report_docx(
        str(path),
        res,
        {
            "tool": "DEM Inundation Boundary",
            "version": "0.3.0",
            "dem": {"path": "dem.tif", "crs": "EPSG:32735", "cell_x": 1, "cell_y": 1},
            "levels_text": "99.500 to 103.000 m",
            "bounding": "box (1 feature(s))",
            "seeds": {"layer": "seed"},
            "outputs": [("Summary per level (CSV)", "x.csv")],
            "warnings": res.warnings,
        },
        geometry,
        bounding=bound,
        seeds_xy=[(cx, cy)],
    )
    doc = docx.Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "DEM Inundation Boundary Assessment" in text
    assert "103.000 m" in text
    assert len(doc.inline_shapes) == 3  # plan, level-area-volume, depth classes
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert not any("," in c and c.replace(",", "").isdigit() for c in cells)
    assert any(w for w in res.warnings if "nothing is inundated" in w)
