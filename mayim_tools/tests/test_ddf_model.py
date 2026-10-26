"""Tests for the duration-consistent DDF model (rfa/ddf_model.py) and its
use in run_frequency_analysis (Weiss factor, time-zone offset, whole
multiples, bootstrap bounds, monotonicity)."""

import numpy as np
import pandas as pd
import pytest

from mayim_tools.rainfall._common.rfa.analysis import run_frequency_analysis
from mayim_tools.rainfall._common.rfa.ddf_model import (
    _pava,
    duration_scale,
    enforce_consistency,
    fit_ddf_model,
)
from mayim_tools.rainfall._common.rfa.distributions import (
    fit_gev_lmoments,
    quantile_for,
)
from mayim_tools.rainfall._common.rfa.timebase import whole_multiple_durations

RPS = [2, 5, 10, 20, 50, 100, 200]


def _scale_invariant_ams(n=60, seed=0, theta=20.0, eta=0.7, cv=0.25, t3=0.2):
    rng = np.random.default_rng(seed)
    xi, al, ka = fit_gev_lmoments(1.0, cv, t3)
    z = np.array([quantile_for("GEV", u, xi, al, ka) for u in rng.random(n)])
    return {
        d: 30.0 * duration_scale(d, theta, eta) * z for d in (60, 120, 360, 720, 1440)
    }


def _hourly_df(years=30, seed=1):
    rng = np.random.default_rng(seed)
    idx = pd.date_range("1990-01-01", periods=years * 8760, freq="h")
    wet = rng.random(len(idx)) < 0.05
    x = np.where(wet, rng.gamma(0.6, 4.0, len(idx)), 0.0)
    # a few storms per year with persistence
    for s in rng.integers(0, len(idx) - 48, years * 6):
        x[s : s + rng.integers(3, 30)] += rng.gamma(2.0, 3.0)
    return pd.DataFrame({"Time": idx.astype(str), "Depth": x})


def test_duration_scale_increasing():
    d = np.geomspace(5, 10080, 100)
    for th, et in ((0, 0.9), (720, 0.0), (60, 0.5)):
        assert np.all(np.diff(duration_scale(d, th, et)) > 0)


def test_parameter_recovery_and_monotone():
    m = fit_ddf_model(_scale_invariant_ams(n=400), vary_cv=False)
    assert m.lambda1 == pytest.approx(30.0 * 1.0, rel=0.08)
    assert m.cv60 == pytest.approx(0.25, abs=0.03)
    assert m.tau3 == pytest.approx(0.2, abs=0.05)
    assert m.is_monotone(m.durations, RPS)


def test_smoothed_growth_follows_duration_trend():
    ams = _scale_invariant_ams(n=300)
    # make short durations more variable: stretch the 60/120-min maxima
    for d, f in ((60, 1.6), (120, 1.3)):
        x = ams[d]
        ams[d] = x.mean() + f * (x - x.mean())
    m0 = fit_ddf_model(ams, vary_cv=False)
    m1 = fit_ddf_model(ams, vary_cv=True)
    assert m1.beta < 0
    lm = m1.per_duration[60.0]
    xi, al, ka = fit_gev_lmoments(lm["l1"], lm["l2"], lm["t3"])
    ind = quantile_for("GEV", 0.99, xi, al, ka)
    assert abs(m1.depth(60, 100) - ind) < abs(m0.depth(60, 100) - ind)


def test_pava_and_enforce_consistency():
    assert list(_pava([1, 3, 2, 4])) == [1, 2.5, 2.5, 4]
    table = {60.0: [10, 20, 30], 120.0: [12, 19, 35], 240.0: [15, 25, 34]}
    new, change = enforce_consistency(table, list(table), [2, 10, 100])
    arr = np.array([new[d] for d in sorted(new)])
    assert np.all(np.diff(arr, axis=0) >= -1e-9)
    assert np.all(np.diff(arr, axis=1) >= -1e-9)
    assert change > 0


def test_too_few_durations_raises():
    ams = _scale_invariant_ams()
    with pytest.raises(ValueError):
        fit_ddf_model({60: ams[60], 120: ams[120]})


def test_whole_multiples():
    use, skip = whole_multiple_durations(60, (5, 30, 60, 90, 120, 1440))
    assert list(use) == [60, 120, 1440]
    assert 90 in skip


def test_run_analysis_ddf_consistent_with_bounds_weiss_tz():
    df = _hourly_df()
    kw = dict(
        requested_durations_min=(60, 90, 120, 180, 360, 720, 1440, 2880),
        exceedance_probabilities=(0.5, 0.1, 0.02, 0.01, 0.005),
        whole_multiples_only=True,
    )
    r = run_frequency_analysis(
        df,
        "Time",
        "Depth",
        fixed_interval_correction=True,
        timezone_offset_h=2,
        n_bootstrap=60,
        **kw,
    )
    assert 90.0 not in r.ddf_model.durations
    assert any("not whole multiples" in w for w in r.warnings)
    ds60 = next(d for d in r.duration_series if d.duration_minutes == 60)
    assert ds60.fixed_interval_factor == pytest.approx(1 / (1 - 1 / 8))
    assert r.metadata["timezone_offset_h"] == 2
    rows = pd.DataFrame(r.ddf_rows)
    piv = rows.pivot(
        index="duration_min", columns="return_period_yr", values="depth_mm"
    )
    assert np.all(np.diff(piv.values, axis=0) >= -1e-9)
    assert np.all(np.diff(piv.values, axis=1) >= -1e-9)
    assert np.all(rows["lower_mm"] <= rows["depth_mm"] + 1e-6)
    assert np.all(rows["upper_mm"] >= rows["depth_mm"] - 1e-6)
    assert r.model_comparison

    r0 = run_frequency_analysis(
        df, "Time", "Depth", fixed_interval_correction=False, **kw
    )
    d0 = next(d for d in r0.duration_series if d.duration_minutes == 60)
    assert d0.fixed_interval_factor == 1.0
    assert r.ddf_rows[0]["depth_mm"] > r0.ddf_rows[0]["depth_mm"]


def test_report_writes(tmp_path):
    pytest.importorskip("docx")
    from mayim_tools.rainfall.frequency_analysis import report

    r = run_frequency_analysis(
        _hourly_df(years=15),
        "Time",
        "Depth",
        requested_durations_min=(60, 120, 360, 1440),
        n_bootstrap=20,
        whole_multiples_only=True,
    )
    out = tmp_path / "r.docx"
    report.write_docx(r, out, inputs={"series": "x.csv"})
    assert out.stat().st_size > 10_000


# ---------------------------------------------------------------- output standard


@pytest.fixture(scope="module")
def small_result():
    return run_frequency_analysis(
        _hourly_df(years=15),
        "Time",
        "Depth",
        requested_durations_min=(60, 120, 360, 1440),
        n_bootstrap=20,
        whole_multiples_only=True,
    )


def test_csvs_follow_the_output_standard(small_result, tmp_path):
    """UTF-8 with BOM everywhere; diagnostic CSVs start with a '#' column
    key covering every heading; the recommended DDF has no key and still
    reads as a DDF table."""
    from mayim_tools.rainfall._common import ddf as ddf_common
    from mayim_tools.rainfall._common.rfa import export as ex

    keyed = {
        "params": (ex.write_parameters, ex.PARAMETERS_KEY),
        "quant": (ex.write_quantiles, ex.QUANTILES_KEY),
        "recs": (ex.write_recommendations, ex.RECOMMENDATIONS_KEY),
        "ams": (ex.write_ams, ex.AMS_KEY),
        "years": (ex.write_year_completeness, ex.YEAR_COMPLETENESS_KEY),
    }
    for name, (fn, key) in keyed.items():
        path = tmp_path / f"{name}.csv"
        fn(small_result, path)
        raw = path.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf"), name
        lines = raw.decode("utf-8-sig").splitlines()
        assert lines[0] == "# Column key", name
        header = next(ln for ln in lines if not ln.startswith("#")).split(",")
        keyed_cols = {ln[2:].split(":")[0] for ln in lines if ln.startswith("# ")}
        patterns = [k for k in key if k.endswith("...")]
        for col in header:
            assert col in keyed_cols or patterns, (name, col)
        df = pd.read_csv(path, comment="#")
        assert len(df) > 0 and df.columns[0] == header[0]

    meta = tmp_path / "meta.csv"
    ex.write_metadata(small_result, meta)
    assert meta.read_bytes().startswith(b"\xef\xbb\xbf")

    ddf = tmp_path / "ddf.csv"
    ex.write_recommended_ddf(small_result, ddf)
    text = ddf.read_text(encoding="utf-8-sig")
    assert ddf.read_bytes().startswith(b"\xef\xbb\xbf")
    assert text.startswith("Duration,")
    long_table = ddf_common.load_ddf_table(str(ddf))
    assert len(long_table) > 0


def test_charts_are_house_size(small_result):
    pytest.importorskip("matplotlib")
    from PIL import Image

    from mayim_tools.rainfall.frequency_analysis import charts

    for fn in (charts.ddf_figure, charts.ratio_figure, charts.lmoment_figure):
        stream, _ = fn(small_result)
        assert Image.open(stream).size[0] == round(16.5 / 2.54 * 200)


def test_report_follows_the_report_standard(small_result, tmp_path):
    pytest.importorskip("docx")
    import re

    from docx import Document

    from mayim_tools.rainfall.frequency_analysis import report

    out = tmp_path / "r.docx"
    report.write_docx(
        small_result,
        out,
        inputs={"series": "x.csv", "site": "BH01"},
        context={
            "project_number": "MC2025_014",
            "project_name": "Riverside Estate",
            "site": "BH01",
            "tool_version": "2.0.0",
            "outputs": [("Recommended DDF (CSV)", "D:/p/MC2025_014_DDF.csv")],
        },
    )
    doc = Document(str(out))
    assert doc.paragraphs[0].text == report.TITLE
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads == [
        "Summary",
        "Tool and method",
        "Inputs and settings",
        "Parameters and calculations",
        "Results",
        "Outputs",
        "Notes and limitations",
        "References",
    ]
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "MC2025_014" in cells and "Riverside Estate" in cells
    assert "Figure 1:" in text and "Table 1:" in text
    assert not re.search(r"\d,\d{3}", text + cells)
    assert "1.0101-year" not in cells
