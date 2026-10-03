"""
Word (.docx) report and figure helpers shared by the Hydrological Tools
reports - zero QGIS dependency. python-docx and matplotlib are imported
lazily inside the functions, so tools still work without them. Figures
use matplotlib's object API (no pyplot), which is safe on QGIS's
Processing worker thread.

Moved verbatim from hydrology/catchment_delineation/report.py so that
Catchment Delineation and DEM depression stage-storage share one
implementation (names made public; that module imports them under its
old private names).
"""

from __future__ import annotations

import io

# Chart styling - a validated categorical palette, recessive grid and ink.
C_BLUE = "#2a78d6"
C_ORANGE = "#eb6834"
C_AQUA = "#1baf7a"
C_INK = "#0b0b0b"
C_INK_2 = "#52514e"
C_GRID = "#e4e3df"
FIG_WIDTH_CM = 16.0


def format_value(value, fmt, scale) -> str:
    if value is None or value == "":
        return "-"
    if scale is not None:
        value = value * scale
    try:
        return fmt.format(value)
    except (ValueError, TypeError):
        return str(value)


def new_figure(height_cm: float):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=(FIG_WIDTH_CM / 2.54, height_cm / 2.54), dpi=200)
    FigureCanvasAgg(fig)
    return fig


def style_axes(ax):
    ax.grid(True, color=C_GRID, linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(C_INK_2)
        ax.spines[side].set_linewidth(0.6)
    ax.tick_params(colors=C_INK_2, labelsize=7, width=0.6)
    ax.xaxis.label.set_color(C_INK)
    ax.yaxis.label.set_color(C_INK)
    ax.xaxis.label.set_size(8)
    ax.yaxis.label.set_size(8)


def png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    buf.seek(0)
    return buf


class Doc:
    """Thin wrapper keeping table/figure numbering."""

    def __init__(self):
        from docx import Document
        from docx.shared import Pt

        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(10)
        self.n_table = 0
        self.n_figure = 0

    def heading(self, text, level):
        self.doc.add_heading(text, level=level)

    def para(self, text, bold_lead=None, style=None):
        p = self.doc.add_paragraph(style=style)
        if bold_lead:
            p.add_run(bold_lead + " ").bold = True
        p.add_run(text)
        return p

    def label(self, text):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = True
        return p

    def bullet(self, text):
        return self.para(text, style="List Bullet")

    def numbered(self, text):
        return self.para(text, style="List Number")

    def caption(self, kind, text):
        from docx.shared import Pt

        if kind == "Table":
            self.n_table += 1
            n = self.n_table
        else:
            self.n_figure += 1
            n = self.n_figure
        p = self.doc.add_paragraph()
        run = p.add_run(f"{kind} {n}: {text}")
        run.italic = True
        run.font.size = Pt(9)
        return p

    def table(self, header, rows, widths_cm=None, group_rows=()):
        """header: list[str]; rows: list[list[str]]; group_rows: indices of
        rows rendered as bold shaded sub-headings spanning the table."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        def shade(cell, fill):
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tc_pr.append(shd)

        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            shade(cell, "D9E2F3")
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            if r_i in group_rows:
                merged = cells[0].merge(cells[-1])
                merged.text = ""
                run = merged.paragraphs[0].add_run(row[0])
                run.bold = True
                run.font.size = Pt(9)
                shade(merged, "F2F2F2")
                continue
            for i, value in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(value))
                run.font.size = Pt(9)
        if widths_cm:
            for i, w in enumerate(widths_cm):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    if i < len(row.cells):
                        row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def picture(self, stream, caption):
        from docx.shared import Cm

        self.doc.add_picture(stream, width=Cm(FIG_WIDTH_CM))
        self.caption("Figure", caption)
