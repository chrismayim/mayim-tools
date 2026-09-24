"""Word (.docx) annexure report for the catchment delineation tool.

Zero QGIS dependency. Needs python-docx and matplotlib (both imported
lazily, so the rest of the tool works without them). Figures are drawn
with matplotlib's object API (no pyplot), which is safe on QGIS's
Processing worker thread.
"""

from __future__ import annotations

import io
from collections.abc import Sequence

import numpy as np

from .core import CatchmentResult
from .description import REFERENCES, describe_catchment

REPORT_TITLE = "Catchment Delineation and Characterisation"

# Chart styling - a validated categorical palette, recessive grid and ink.
C_BLUE = "#2a78d6"
C_ORANGE = "#eb6834"
C_AQUA = "#1baf7a"
C_INK = "#0b0b0b"
C_INK_2 = "#52514e"
C_GRID = "#e4e3df"
FIG_WIDTH_CM = 16.0

# (parameter, symbol and unit, definition / equation, source)
PARAMETER_DEFINITIONS = [
    ("Contributing area", "A (km²)", "Number of cells draining to the outlet x cell area.", "-"),
    ("Polygon area", "A_p (km²)", "Area of the catchment outline, including filled interior holes.", "-"),
    ("Perimeter", "P (km)", "Length of the catchment outline, following cell edges.", "-"),
    ("Elevations", "z (m)", "Minimum, maximum and mean DEM elevation of contributing cells; z_o at the outlet cell.", "-"),
    ("Longest flow path", "L (km)", "Longest D8 flow path to the outlet, from cell centre to outlet cell centre.", "O'Callaghan & Mark (1984)"),
    ("Centroid distance", "Lca (km)", "Distance along L from the outlet to the point on L nearest the catchment centroid.", "-"),
    ("Basin length", "Lb (km)", "Longest straight-line distance from the outlet to any point in the catchment.", "-"),
    ("Average slope", "S_avg (m/m)", "S_avg = (z_top - z_o) / L, along the longest flow path.", "-"),
    ("10-85 slope", "S_1085 (m/m)", "S_1085 = (z_85 - z_10) / (0.75 L), elevations at 10 % and 85 % of L from the outlet.", "-"),
    ("Equal-area slope", "S_e (m/m)", "Slope of the line through the outlet with the same area under it as the profile: S_e = 2 ∫(z - z_o) dx / L².", "-"),
    ("Mean catchment slope", "S_m (%)", "Mean of the cell slopes (3 x 3 finite-difference method) over the catchment.", "Horn (1981)"),
    ("Form factor", "R_f", "R_f = A_p / L².", "Horton (1932)"),
    ("Circularity ratio", "R_c", "R_c = 4π A_p / P².", "Miller (1953)"),
    ("Elongation ratio", "R_e", "R_e = 2 √(A_p / π) / L.", "Schumm (1956)"),
    ("Compactness coefficient", "K_c", "K_c = P / (2 √(π A_p)).", "Gravelius (1914)"),
    ("Hypsometric integral", "HI", "HI = (z_mean - z_min) / (z_max - z_min) (elevation-relief ratio).", "Strahler (1952); Pike & Wilson (1971)"),
    ("Relief ratio", "R_h", "R_h = H / Lb, with H = z_max - z_o.", "Schumm (1956)"),
    ("Melton ratio", "M", "M = H / √A (H in km).", "Melton (1965)"),
    ("Ruggedness number", "R_n", "R_n = H x D_d (H in km).", "Strahler (1958)"),
    ("Streams", "-", "Cells whose upstream area is at least the stream threshold.", "-"),
    ("Stream order", "u", "Strahler ordering of the stream cells; the catchment order is the highest.", "Strahler (1957)"),
    ("Bifurcation ratio", "R_b", "Mean of N_u / N_u+1 over consecutive orders (N_u = number of segments of order u).", "Strahler (1964)"),
    ("Drainage density", "D_d (km/km²)", "D_d = total stream length / A.", "Horton (1932)"),
    ("Stream frequency", "F_s (1/km²)", "F_s = number of stream segments / A.", "Horton (1932)"),
    ("Length of overland flow", "L_o (km)", "L_o = 1 / (2 D_d).", "Horton (1945)"),
    ("Constant of channel maintenance", "C (km²/km)", "C = 1 / D_d.", "Schumm (1956)"),
    ("Sinuosity", "-", "L / straight-line distance between its ends.", "-"),
    ("Tc - Kirpich", "t_c (min)", "t_c = 0.0195 L^0.77 S_avg^-0.385 (L in m).", "Kirpich (1940)"),
    ("Tc - USBR", "t_c (min)", "t_c = 60 (0.87 L² / (1000 S_1085))^0.385 (L in km).", "USBR (1973)"),
    ("Tc - Bransby-Williams", "t_c (min)", "t_c = 58 L / (A^0.1 S_e^0.2) (L in km, A in km², S_e in m/km).", "Bransby-Williams (1922)"),
]  # fmt: skip

# (group heading, [(label, attribute key, format, scale)])
RESULT_GROUPS = [
    (
        "Geometry",
        [
            ("Contributing area A (km²)", "area_km2", "{:.4f}", 1),
            ("Contributing area (ha)", "area_ha", "{:.2f}", 1),
            ("Polygon area A_p (km²)", "poly_km2", "{:.4f}", 1),
            ("Filled interior holes (ha)", "hole_ha", "{:.2f}", 1),
            ("Perimeter P (km)", "perim_km", "{:.3f}", 1),
            ("Polygon parts", "n_parts", "{:d}", None),
            ("Centroid E", "cx", "{:.1f}", 1),
            ("Centroid N", "cy", "{:.1f}", 1),
        ],
    ),
    (
        "Outlet",
        [
            ("Outlet type", "out_type", "{}", None),
            ("Outlet E", "x_out", "{:.1f}", 1),
            ("Outlet N", "y_out", "{:.1f}", 1),
            ("Snap distance (m)", "snap_m", "{:.1f}", 1),
            ("Outlet cells", "n_outcells", "{:d}", None),
            ("Drains into", "ds_ids", "{}", None),
        ],
    ),
    (
        "Elevation and relief",
        [
            ("Outlet elevation z_o (m)", "z_outlet", "{:.2f}", 1),
            ("Minimum elevation (m)", "z_min", "{:.2f}", 1),
            ("Maximum elevation (m)", "z_max", "{:.2f}", 1),
            ("Mean elevation (m)", "z_mean", "{:.2f}", 1),
            ("Relief z_max - z_min (m)", "relief_m", "{:.2f}", 1),
            ("Hypsometric integral HI", "hyps_int", "{:.3f}", 1),
            ("Relief ratio R_h", "relief_r", "{:.4f}", 1),
            ("Melton ratio M", "melton_r", "{:.3f}", 1),
            ("Ruggedness number R_n", "rugged_n", "{:.3f}", 1),
        ],
    ),
    (
        "Lengths and slopes",
        [
            ("Longest flow path L (km)", "lfp_km", "{:.3f}", 1),
            ("Centroid distance Lca (km)", "lca_km", "{:.3f}", 1),
            ("Basin length Lb (km)", "lb_km", "{:.3f}", 1),
            ("Sinuosity of L", "sinuosity", "{:.3f}", 1),
            ("Average slope S_avg (%)", "s_avg", "{:.3f}", 100),
            ("10-85 slope S_1085 (%)", "s_1085", "{:.3f}", 100),
            ("Equal-area slope S_e (%)", "s_ea", "{:.3f}", 100),
            ("Mean catchment slope S_m (%)", "s_mean_pc", "{:.2f}", 1),
        ],
    ),
    (
        "Shape",
        [
            ("Form factor R_f", "form_f", "{:.3f}", 1),
            ("Circularity ratio R_c", "circ_r", "{:.3f}", 1),
            ("Elongation ratio R_e", "elong_r", "{:.3f}", 1),
            ("Compactness coefficient K_c", "gravel_kc", "{:.3f}", 1),
        ],
    ),
    (
        "Drainage network",
        [
            ("Highest Strahler order", "strm_ord", "{:d}", None),
            ("Stream segments", "n_strm", "{:d}", None),
            ("Stream length (km)", "strm_km", "{:.3f}", 1),
            ("Drainage density D_d (km/km²)", "dd_kmkm2", "{:.3f}", 1),
            ("Stream frequency F_s (1/km²)", "strm_freq", "{:.3f}", 1),
            ("Bifurcation ratio R_b", "bif_r", "{:.2f}", 1),
            ("Length of overland flow L_o (km)", "lo_km", "{:.3f}", 1),
            ("Constant of channel maintenance C (km²/km)", "c_maint", "{:.3f}", 1),
        ],
    ),
    (
        "Time of concentration",
        [
            ("Kirpich (min)", "tc_kirp_mn", "{:.1f}", 1),
            ("USBR (min)", "tc_usbr_mn", "{:.1f}", 1),
            ("Bransby-Williams (min)", "tc_bw_mn", "{:.1f}", 1),
        ],
    ),
]


def format_value(value, fmt, scale) -> str:
    if value is None or value == "":
        return "-"
    if scale is not None:
        value = value * scale
    try:
        return fmt.format(value)
    except (ValueError, TypeError):
        return str(value)


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------


def _new_figure(height_cm: float):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=(FIG_WIDTH_CM / 2.54, height_cm / 2.54), dpi=200)
    FigureCanvasAgg(fig)
    return fig


def _style_axes(ax):
    ax.grid(True, color=C_GRID, linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(C_INK_2)
        ax.spines[side].set_linewidth(0.6)
    ax.tick_params(colors=C_INK_2, labelsize=7, width=0.6)
    ax.xaxis.label.set_color(C_INK)
    ax.yaxis.label.set_color(C_INK)
    ax.xaxis.label.set_size(8)
    ax.yaxis.label.set_size(8)


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    buf.seek(0)
    return buf


def figure_plan(c: CatchmentResult) -> io.BytesIO:
    from matplotlib.collections import LineCollection
    from matplotlib.patches import PathPatch
    from matplotlib.path import Path

    fig = _new_figure(13.0)
    ax = fig.add_subplot(1, 1, 1)
    for poly in c.polygons:
        verts, codes = [], []
        for ring in poly:
            verts.extend(ring)
            codes.extend([Path.MOVETO] + [Path.LINETO] * (len(ring) - 2))
            codes.append(Path.CLOSEPOLY)
        ax.add_patch(
            PathPatch(
                Path(verts, codes),
                facecolor=C_BLUE,
                alpha=0.10,
                edgecolor="none",
            )
        )
        ax.add_patch(
            PathPatch(
                Path(verts, codes), facecolor="none", edgecolor=C_INK, linewidth=0.8
            )
        )

    streams = c.detail.get("streams")
    if streams is not None and len(streams):
        orders = streams[:, 4]
        max_order = max(orders.max(), 1)
        segs = streams[:, :4].reshape(-1, 2, 2)
        widths = 0.4 + 1.4 * (orders - 1) / max(max_order - 1, 1)
        ax.add_collection(
            LineCollection(segs, colors=C_BLUE, linewidths=widths, capstyle="round")
        )
        ax.plot([], [], color=C_BLUE, linewidth=1.2, label="Streams (width by order)")

    if len(c.flow_path) > 1:
        fx, fy = zip(*c.flow_path, strict=True)
        ax.plot(fx, fy, color=C_ORANGE, linewidth=1.6, label="Longest flow path")
    if c.outlet_xy:
        ax.plot(
            *c.outlet_xy,
            marker="v",
            markersize=8,
            color=C_INK,
            linestyle="none",
            label="Outlet",
        )
    cen = c.detail.get("centroid")
    if cen:
        ax.plot(
            *cen,
            marker="+",
            markersize=9,
            markeredgewidth=1.5,
            color=C_INK_2,
            linestyle="none",
            label="Centroid",
        )
    ax.set_aspect("equal", adjustable="datalim")
    ax.autoscale_view()
    ax.set_xlabel("Easting (m)")
    ax.set_ylabel("Northing (m)")
    _style_axes(ax)
    ax.ticklabel_format(useOffset=False, style="plain")
    ax.legend(fontsize=7, frameon=False, loc="best")
    return _png(fig)


def figure_profile(c: CatchmentResult) -> io.BytesIO | None:
    x = np.asarray(c.detail.get("profile_x_m", []), dtype=float)
    z = np.asarray(c.detail.get("profile_z_m", []), dtype=float)
    if x.size < 2 or np.isnan(z).all():
        return None
    a = c.attributes
    fig = _new_figure(7.5)
    ax = fig.add_subplot(1, 1, 1)
    xk = x / 1000.0
    ax.plot(xk, z, color=C_BLUE, linewidth=1.6, label="Longest flow path profile")
    length = x[-1]
    if a.get("s_1085") is not None:
        x10, x85 = 0.10 * length, 0.85 * length
        z10 = float(np.interp(x10, x, z))
        z85 = float(np.interp(x85, x, z))
        ax.plot(
            [x10 / 1000, x85 / 1000],
            [z10, z85],
            color=C_ORANGE,
            linewidth=1.2,
            linestyle="--",
            label=f"10-85 slope {100 * a['s_1085']:.2f} %",
        )
        ax.plot([x10 / 1000, x85 / 1000], [z10, z85], "o", color=C_ORANGE, ms=4)
    if a.get("s_ea") is not None:
        ax.plot(
            [0, length / 1000],
            [z[0], z[0] + a["s_ea"] * length],
            color=C_AQUA,
            linewidth=1.2,
            linestyle=":",
            label=f"Equal-area slope {100 * a['s_ea']:.2f} %",
        )
    ax.set_xlabel("Distance upstream from outlet (km)")
    ax.set_ylabel("Elevation (m)")
    ax.set_xlim(0, xk[-1])
    _style_axes(ax)
    ax.legend(fontsize=7, frameon=False, loc="upper left")
    return _png(fig)


def figure_hypsometric(c: CatchmentResult) -> io.BytesIO | None:
    curve = c.hypsometric_curve
    if not curve:
        return None
    rel_a = [p["rel_area"] for p in curve]
    rel_h = [p["rel_height"] for p in curve]
    fig = _new_figure(7.5)
    ax = fig.add_subplot(1, 1, 1)
    ax.fill_between(rel_a, rel_h, color=C_BLUE, alpha=0.12, linewidth=0)
    ax.plot(rel_a, rel_h, color=C_BLUE, linewidth=1.6)
    ax.plot([0, 1], [1, 0], color=C_INK_2, linewidth=0.6, linestyle="--")
    hi = c.attributes.get("hyps_int")
    if hi is not None:
        ax.text(
            0.97,
            0.93,
            f"HI = {hi:.3f}",
            transform=ax.transAxes,
            ha="right",
            va="top",
            fontsize=8,
            color=C_INK,
        )
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_xlabel("Relative area above height, a/A")
    ax.set_ylabel("Relative height, h/H")
    _style_axes(ax)
    return _png(fig)


def figure_histogram(c: CatchmentResult) -> io.BytesIO | None:
    counts = np.asarray(c.detail.get("elev_hist_counts", []), dtype=float)
    edges = np.asarray(c.detail.get("elev_hist_edges", []), dtype=float)
    if counts.size == 0 or counts.sum() == 0:
        return None
    pct = 100.0 * counts / counts.sum()
    fig = _new_figure(7.0)
    ax = fig.add_subplot(1, 1, 1)
    ax.bar(
        edges[:-1],
        pct,
        width=np.diff(edges),
        align="edge",
        color=C_BLUE,
        edgecolor="white",
        linewidth=1.0,
    )
    ax.set_xlabel("Elevation (m)")
    ax.set_ylabel("Share of catchment area (%)")
    _style_axes(ax)
    return _png(fig)


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------


class _Doc:
    """Thin wrapper keeping table/figure numbering."""

    def __init__(self):
        from docx import Document
        from docx.shared import Pt

        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(10)
        self.n_table = 0
        self.n_figure = 0

    def heading(self, text, level):
        self.doc.add_heading(text, level=level)

    def para(self, text, bold_lead=None, style=None):
        p = self.doc.add_paragraph(style=style)
        if bold_lead:
            p.add_run(bold_lead + " ").bold = True
        p.add_run(text)
        return p

    def label(self, text):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = True
        return p

    def bullet(self, text):
        return self.para(text, style="List Bullet")

    def numbered(self, text):
        return self.para(text, style="List Number")

    def caption(self, kind, text):
        from docx.shared import Pt

        if kind == "Table":
            self.n_table += 1
            n = self.n_table
        else:
            self.n_figure += 1
            n = self.n_figure
        p = self.doc.add_paragraph()
        run = p.add_run(f"{kind} {n}: {text}")
        run.italic = True
        run.font.size = Pt(9)
        return p

    def table(self, header, rows, widths_cm=None, group_rows=()):
        """header: list[str]; rows: list[list[str]]; group_rows: indices of
        rows rendered as bold shaded sub-headings spanning the table."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        def shade(cell, fill):
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tc_pr.append(shd)

        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            shade(cell, "D9E2F3")
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            if r_i in group_rows:
                merged = cells[0].merge(cells[-1])
                merged.text = ""
                run = merged.paragraphs[0].add_run(row[0])
                run.bold = True
                run.font.size = Pt(9)
                shade(merged, "F2F2F2")
                continue
            for i, value in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(value))
                run.font.size = Pt(9)
        if widths_cm:
            for i, w in enumerate(widths_cm):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    if i < len(row.cells):
                        row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def picture(self, stream, caption):
        from docx.shared import Cm

        self.doc.add_picture(stream, width=Cm(FIG_WIDTH_CM))
        self.caption("Figure", caption)


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def _tc_range(a):
    vals = [a.get(k) for k in ("tc_kirp_mn", "tc_usbr_mn", "tc_bw_mn")]
    vals = [v for v in vals if v is not None]
    if not vals:
        return "-"
    return f"{min(vals):.1f} - {max(vals):.1f}"


def write_report_docx(
    path: str,
    catchments: Sequence[CatchmentResult],
    context: dict,
    include_figures: bool = True,
) -> None:
    """Writes the annexure report. ``context`` describes the run (see the
    QGIS wrapper): tool/version/run time, inputs, settings, outputs and
    processing warnings."""
    d = _Doc()
    ctx = context or {}
    if include_figures:
        try:
            import matplotlib  # noqa: F401
        except ImportError:
            include_figures = False
    dem = ctx.get("dem", {})
    mode = ctx.get("mode", "total")

    d.heading(REPORT_TITLE, 1)

    # 1. Summary ---------------------------------------------------------------
    d.heading("1. Summary", 2)
    n = len(catchments)
    d.para(
        f"This annexure documents the delineation and characterisation of "
        f"{n} catchment{'s' if n != 1 else ''} from a digital elevation model "
        f"(DEM) with a {dem.get('cell_x', 0):g} m x {dem.get('cell_y', 0):g} m "
        f"cell size, using the {ctx.get('tool', 'Catchment Delineation')} tool "
        f"(Mayim Tools, version {ctx.get('version', '-')}) in "
        f"{'total' if mode == 'total' else 'incremental'} mode. For each "
        "catchment it records the inputs and method, the geometric, relief, "
        "shape, drainage-network and time-of-concentration parameters, a "
        "geomorphological description and general considerations for "
        "hydrological assessment and modelling."
    )
    d.caption("Table", "Summary of delineated catchments")
    d.table(
        [
            "Catchment",
            "Area (km²)",
            "L (km)",
            "S_1085 (%)",
            "S_m (%)",
            "Order",
            "Tc range (min)",
        ],
        [
            [
                c.outlet_id,
                format_value(c.attributes["area_km2"], "{:.3f}", 1),
                format_value(c.attributes["lfp_km"], "{:.3f}", 1),
                format_value(c.attributes["s_1085"], "{:.2f}", 100),
                format_value(c.attributes["s_mean_pc"], "{:.1f}", 1),
                format_value(c.attributes["strm_ord"], "{:d}", None),
                _tc_range(c.attributes),
            ]
            for c in catchments
        ],
    )

    # 2. Tool and method -------------------------------------------------------
    d.heading("2. Tool and method", 2)
    d.para(
        "The catchments were delineated with a D8 (eight-direction) flow model "
        "(O'Callaghan & Mark 1984), in which every DEM cell drains to the "
        "steepest-descent neighbour of its eight neighbours. The processing "
        "steps were:"
    )
    for step in [
        "The D8 flow-direction (pointer) raster was decoded and checked; codes "
        "that are not valid D8 directions were treated as NoData and reported.",
        "Flow accumulation (number of upstream cells) was computed in "
        "topological order from the pointer.",
        "Pour points were moved to the cell with the highest flow accumulation "
        f"within {ctx.get('snap_radius_cells', '-')} cells, so that they sit "
        "on the drainage line. Pour lines were rasterised as a 4-connected "
        "chain of cells, so no D8 flow path can cross a line without being "
        "counted.",
        "Every cell was traced downstream to the first outlet cell it reaches. "
        + (
            "In total mode each catchment includes everything upstream of its "
            "outlet, so nested catchments overlap."
            if mode == "total"
            else "In incremental mode each cell belongs only to the first outlet "
            "it meets, which gives non-overlapping sub-catchments linked to the "
            "outlet they drain into."
        ),
        "The outline was traced along cell edges. Interior holes (cells inside "
        "the outline that do not drain to the outlet) were filled and their "
        "area reported. Cells touching only at a corner form separate polygon "
        "parts, and the geometry was checked and repaired to a valid "
        "(multi)polygon.",
        "The DEM was clipped to the catchment outlines, and the parameters, "
        "longest flow path, stream network and figures in this annexure were "
        "derived from the DEM and flow network.",
    ]:
        d.numbered(step)

    # 3. Inputs ---------------------------------------------------------------
    d.heading("3. Inputs and settings", 2)
    d.caption("Table", "Inputs and processing settings")
    d.table(
        ["Item", "Value"],
        [
            ["DEM", dem.get("path", "-")],
            ["Coordinate reference system", dem.get("crs", "-")],
            [
                "DEM cell size",
                f"{dem.get('cell_x', 0):g} m x {dem.get('cell_y', 0):g} m",
            ],
            [
                "DEM dimensions",
                f"{dem.get('cols', '-')} columns x {dem.get('rows', '-')} rows",
            ],
            ["D8 pointer", ctx.get("pointer", {}).get("path", "-")],
            ["Pointer encoding", ctx.get("pointer", {}).get("encoding", "-")],
            ["Outlet layer", ctx.get("outlets", {}).get("layer", "-")],
            [
                "Outlet features",
                str(ctx.get("outlets", {}).get("n_features", "-")),
            ],
            ["Outlet name field", ctx.get("outlets", {}).get("id_field") or "-"],
            ["Pour point snap radius", f"{ctx.get('snap_radius_cells', '-')} cells"],
            ["Catchment mode", "Total" if mode == "total" else "Incremental"],
            [
                "Stream threshold",
                f"{ctx.get('stream_threshold_km2', '-')} km² upstream area",
            ],
            ["Run", ctx.get("run_time", "-")],
        ],
        widths_cm=[5.0, 11.0],
    )
    d.para(
        "The DEM is assumed to be hydrologically conditioned (pits filled or "
        "breached, flats resolved) and in a projected coordinate system in "
        "metres. All lengths, areas and slopes are in that system."
    )

    # 4. Parameters ------------------------------------------------------------
    d.heading("4. Parameters and calculations", 2)
    d.caption("Table", "Parameter definitions")
    d.table(
        ["Parameter", "Symbol (unit)", "Definition / equation", "Source"],
        [list(r) for r in PARAMETER_DEFINITIONS],
        widths_cm=[3.4, 2.6, 7.0, 3.0],
    )
    d.para(
        "Kirpich uses S_avg, USBR uses S_1085 and Bransby-Williams uses S_e, "
        "each with the longest flow path L. None of the three separates "
        "overland from channel flow, so they apply the whole path length as "
        "channel-type flow."
    )

    # 5. Results ---------------------------------------------------------------
    d.heading("5. Catchment results", 2)
    for i, c in enumerate(catchments, start=1):
        a = c.attributes
        text = describe_catchment(a, c.detail)
        d.heading(f"5.{i} Catchment {c.outlet_id}", 3)
        d.para(text["overview"])
        d.label("Geomorphological description")
        for paragraph in text["geomorphology"]:
            lead, _, rest = paragraph.partition(". ")
            d.para(rest, bold_lead=lead + ".")
        d.label("Hydrological assessment and modelling considerations")
        for point in text["modelling"]:
            d.bullet(point)

        rows, groups = [], []
        for title, items in RESULT_GROUPS:
            groups.append(len(rows))
            rows.append([title, ""])
            for label, key, fmt, scale in items:
                rows.append([label, format_value(a.get(key), fmt, scale)])
        d.caption("Table", f"Parameters of catchment {c.outlet_id}")
        d.table(["Parameter", "Value"], rows, widths_cm=[10.0, 6.0], group_rows=groups)

        slope_rows = [
            [s["label"], f"{s['pct']:.1f}"] for s in c.detail.get("slope_classes", [])
        ]
        if slope_rows:
            d.caption(
                "Table",
                f"Slope class distribution, catchment {c.outlet_id} "
                "(FAO 2006 classes)",
            )
            d.table(["Slope class", "Area (%)"], slope_rows, widths_cm=[10.0, 6.0])
        aspect_rows = [
            [s["label"], f"{s['pct']:.1f}"] for s in c.detail.get("aspect", [])
        ]
        if aspect_rows:
            d.caption("Table", f"Aspect distribution, catchment {c.outlet_id}")
            d.table(["Aspect", "Area (%)"], aspect_rows, widths_cm=[10.0, 6.0])
        order_rows = [
            [str(o["order"]), str(o["n_segments"]), f"{o['length_km']:.3f}"]
            for o in c.detail.get("stream_orders", [])
        ]
        if order_rows:
            d.caption("Table", f"Stream network by Strahler order, {c.outlet_id}")
            d.table(
                ["Order u", "Segments N_u", "Length (km)"],
                order_rows,
                widths_cm=[5.0, 5.5, 5.5],
            )

        if include_figures:
            d.picture(
                figure_plan(c),
                f"Catchment {c.outlet_id}: outline, stream network, longest flow "
                "path, outlet and centroid",
            )
            for maker, cap in (
                (figure_profile, "longitudinal profile of the longest flow path"),
                (figure_hypsometric, "hypsometric curve"),
                (figure_histogram, "elevation distribution"),
            ):
                stream = maker(c)
                if stream is not None:
                    d.picture(stream, f"Catchment {c.outlet_id}: {cap}")

    # 6. Outputs ----------------------------------------------------------------
    outputs = ctx.get("outputs", [])
    if outputs:
        d.heading("6. Outputs", 2)
        d.caption("Table", "Files produced by the tool")
        d.table(["Output", "Location"], [list(o) for o in outputs], widths_cm=[5, 11])

    # 7. Notes -----------------------------------------------------------------
    d.heading("7. Notes and limitations", 2)
    for note in [
        "Results are only as good as the DEM. Its resolution, vertical "
        "accuracy, vegetation and building artefacts, and hydrological "
        "conditioning all control the flow directions and hence the boundary.",
        "A bare-earth DEM does not contain culverts, bridges, pipes or "
        "diversions. Where these exist, check the boundary against site "
        "information.",
        "D8 routing sends all flow from a cell to one neighbour and cannot "
        "represent flow splitting (e.g. on fans or flat ground).",
        "Stream network parameters depend on the chosen stream threshold and "
        "the DEM resolution.",
        "Time-of-concentration values are empirical estimates for guidance; "
        "the design method and parameters should follow the governing local "
        "guideline.",
    ]:
        d.bullet(note)
    warnings = ctx.get("warnings", [])
    if warnings:
        d.para("Processing messages recorded during this run:")
        for w in warnings:
            d.bullet(w)

    # 8. References --------------------------------------------------------------
    d.heading("8. References", 2)
    for ref in REFERENCES:
        d.para(ref)

    d.doc.save(path)
