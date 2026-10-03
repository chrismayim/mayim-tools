"""Word (.docx) report for Precipitation Data to DDF.

Requires python-docx (soft import) so the analysis and CSV outputs never
depend on it. Charts are produced by charts.py (matplotlib, optional).
"""

from __future__ import annotations

import datetime as _dt
import math
import tempfile
from pathlib import Path

import numpy as np

from . import charts

REFERENCES = [
    "Hosking, J.R.M. & Wallis, J.R. (1997). Regional Frequency Analysis: An "
    "Approach Based on L-Moments. Cambridge University Press.",
    "Koutsoyiannis, D., Kozonis, D. & Manetas, A. (1998). A mathematical framework "
    "for studying rainfall intensity-duration-frequency relationships. Journal of "
    "Hydrology, 206(1-2), 118-135.",
    "Overeem, A., Buishand, A. & Holleman, I. (2008). Rainfall depth-duration-"
    "frequency curves and their uncertainties. Journal of Hydrology, 348(1-2), "
    "124-134.",
    "Roksvag, T., Lutz, J., Grinde, L., Dyrrdal, A.V. & Thorarinsdottir, T.L. "
    "(2021). Consistent intensity-duration-frequency curves by post-processing of "
    "estimated Bayesian posterior quantiles. Journal of Hydrology, 603, 127000.",
    "Smithers, J.C. & Schulze, R.E. (2003). Design Rainfall and Flood Estimation "
    "in South Africa. WRC Report K5/1060, Water Research Commission, Pretoria.",
    "Ball, J., et al. (eds) (2019). Australian Rainfall and Runoff: A Guide to "
    "Flood Estimation, Book 2 (Rainfall Estimation). Commonwealth of Australia "
    "(Geoscience Australia).",
    "Weiss, L.L. (1964). Ratio of true to fixed-interval maximum rainfall. Journal "
    "of the Hydraulics Division, ASCE, 90(1), 77-82.",
    "Vogel, R.M. & Fennessey, N.M. (1993). L moment diagrams should replace product "
    "moment diagrams. Water Resources Research, 29(6), 1745-1752.",
    "Efron, B. & Tibshirani, R.J. (1993). An Introduction to the Bootstrap. "
    "Chapman & Hall.",
    "Barlow, R.E., Bartholomew, D.J., Bremner, J.M. & Brunk, H.D. (1972). "
    "Statistical Inference under Order Restrictions. Wiley.",
]


def _f(v, nd=2):
    if v is None or v == "":
        return "-"
    try:
        fv = float(v)
        if math.isnan(fv) or math.isinf(fv):
            return "-"
        return f"{fv:.{nd}f}"
    except (TypeError, ValueError):
        return str(v)


def _table(doc, header, rows, font_pt=None):
    from docx.shared import Pt

    t = doc.add_table(rows=1, cols=len(header))
    t.style = (
        "Light Grid Accent 1"
        if "Light Grid Accent 1" in [s.name for s in doc.styles]
        else "Table Grid"
    )
    for j, h in enumerate(header):
        t.rows[0].cells[j].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            cells[j].text = str(v)
    if font_pt:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(font_pt)
    return t


def _caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True


def write_docx(result, path, inputs=None) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    inputs = inputs or {}
    md = result.metadata
    m = result.ddf_model
    rps = list(result.return_periods)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    doc.add_heading("Rainfall Depth-Duration-Frequency Analysis", level=0)
    doc.add_paragraph(
        f"Site / series: {inputs.get('site') or Path(inputs.get('series', '-')).stem}"
        f"    Generated: {_dt.date.today().isoformat()}    "
        "Tool: Mayim Tools - Precipitation Data to DDF"
    )

    # 1 Purpose
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "This report documents the derivation of a Depth-Duration-Frequency (DDF) "
        "table from a recorded precipitation time series (gauge, or gridded "
        "reanalysis / satellite product). Annual maximum series (AMS) are extracted "
        "at each duration, frequency distributions are fitted, and a single "
        "duration-consistent DDF model is fitted across all durations so that the "
        "recommended design depths always increase with both duration and return "
        "period. Per-duration fits and the L-moment ratio diagram are reported as "
        "diagnostics."
    )

    # 2 Inputs
    doc.add_heading("2. Inputs and data", level=1)
    ds_list = result.duration_series
    n_years = max((len(d.values_mm) for d in ds_list), default=0)
    yrs = sorted({r.year for d in ds_list for r in d.year_records})
    excl = sorted({r.year for d in ds_list for r in d.year_records if not r.included})
    rows = [
        ("Rainfall series", inputs.get("series", "-")),
        ("Record", f"{md.get('record_start')} to {md.get('record_end')}"),
        ("Native time step", f"{_f(md.get('native_interval_min'), 0)} min"),
        ("Time-zone offset applied", f"{_f(md.get('timezone_offset_h'), 1)} h"),
        ("Year definition", md.get("year_definition", "calendar year")),
        (
            "Years in record / used",
            f"{len(yrs)} / {n_years} (minimum completeness "
            f"{_f(100 * float(md.get('min_completeness', 0.9)), 0)}%)",
        ),
        ("Years excluded", ", ".join(str(y) for y in excl) or "none"),
        (
            "Durations analysed",
            ", ".join(d.duration_label for d in ds_list),
        ),
        ("Fixed-interval correction", md.get("fixed_interval_correction")),
        (
            "Distributions fitted per duration",
            ", ".join(md.get("distributions_fitted", [])),
        ),
        ("GEV fitting method (per duration)", md.get("gev_method")),
        ("DDF model distribution", md.get("ddf_model_distribution", "-")),
        ("DDF model growth curve", md.get("ddf_model_growth", "-")),
        (
            "Bootstrap",
            f"{md.get('n_bootstrap', 0)} replicates (seed {md.get('random_seed')})",
        ),
    ]
    _table(doc, ["Item", "Value"], rows)

    # 3 Method
    doc.add_heading("3. Method", level=1)
    doc.add_paragraph(
        "Annual maximum series. For each duration the maximum depth accumulated in a "
        "sliding window of that length is extracted for each calendar year. Missing "
        "data are never treated as zero; years whose completeness is below the "
        "threshold are excluded from every duration so that all durations share the "
        "same years. Return periods are annual-maximum values (AEP = 1/T)."
    )
    if md.get("fixed_interval_correction") not in (None, "off"):
        doc.add_paragraph(
            "Fixed-interval correction. Maxima from data recorded at a fixed "
            "interval under-estimate the true (unrestricted) maxima because a storm "
            "rarely aligns with the clock. Each window of n native steps was "
            "multiplied by the Weiss (1964) factor 1/(1 - 1/(8n)), e.g. 1.143 for "
            "n = 1, 1.021 for n = 6 and 1.005 for n = 24."
        )
    doc.add_paragraph(
        "Per-duration fits (diagnostic). The GEV, Gumbel, GLO and LP3 distributions "
        "are fitted independently to each duration by L-moments (Hosking & Wallis "
        "1997). The L-moment ratio diagram (Vogel & Fennessey 1993) ranks the "
        "candidates at each duration by the distance between the sample (tau3, tau4) "
        "and each distribution's theoretical curve (Gumbel: a single point, "
        "measured in both tau3 and tau4). Fitting each duration separately - and "
        "especially selecting a different distribution for each duration - can "
        "produce DDF curves that cross, i.e. a longer duration with a smaller "
        "design depth, which is physically impossible. These fits are therefore "
        "reported for review only."
    )
    doc.add_paragraph(
        "Duration-consistent DDF model (recommended). Following the "
        "index-flood-across-durations framework of Koutsoyiannis et al. (1998) and "
        "Overeem et al. (2008), the annual maximum depth X_d at duration d (minutes) "
        "is written as X_d = l1(d) * Y_d, where the mean scales smoothly with "
        "duration, l1(d) = lambda1 * (d/60) * ((d + theta)/(60 + theta))^(-eta) with "
        "theta >= 0 and 0 <= eta < 1 (strictly increasing in d), and Y_d is a growth "
        "variable with mean 1 and a single distribution family (GEV by default). "
        "theta, eta and lambda1 are fitted by record-length-weighted least squares on "
        "the logarithm of the sample means."
    )
    if md.get("ddf_model_growth", "").startswith("smoothed"):
        doc.add_paragraph(
            "Growth curve. Short durations are typically more variable and more "
            "skewed than long durations, so the L-CV and L-skewness of Y_d are "
            "allowed to vary smoothly with duration: L-CV(d) = cv60 * (d/60)^beta "
            "and tau3(d) = tau3_60 + g * ln(d/60), fitted by weighted least squares "
            "to the sample L-moments of all durations. This smooths the at-site "
            "L-moments across durations (cf. the regression smoothing of L-moments "
            "across durations in the ARR 2016 IFDs; Ball et al. 2019) instead of "
            "fitting each duration in isolation, which borrows strength across "
            "durations and damps sampling noise in the shape parameter."
        )
    else:
        doc.add_paragraph(
            "Growth curve. A single growth curve (constant L-CV and L-skewness, "
            "record-length-weighted means across durations) is used for all "
            "durations - the simple scaling (scale-invariance) assumption."
        )
    doc.add_paragraph(
        "Consistency. The model table is checked for depths increasing with duration "
        "and return period; any remaining crossings (typically only possible at the "
        "longest durations or rarest return periods, where the fitted L-skewness "
        "trend is extrapolated) are removed by isotonic regression "
        "(pool-adjacent-violators; Barlow et al. 1972), as proposed for IDF curves "
        "by Roksvag et al. (2021). The size of any such correction is reported."
    )
    if int(md.get("n_bootstrap") or 0):
        doc.add_paragraph(
            f"Uncertainty. 5-95% bounds are from {md.get('n_bootstrap')} "
            "non-parametric bootstrap replicates (Efron & Tibshirani 1993) in which "
            "whole years are resampled with replacement - the same years for every "
            "duration, so the dependence between durations is preserved - and the "
            "full model (scaling, growth curve and consistency step) is refitted "
            "each time. The bounds express sampling uncertainty only, not "
            "uncertainty in the data themselves or in the choice of model."
        )

    # 4 Model parameters
    if m is not None:
        doc.add_heading("4. DDF model parameters", level=1)
        prow = [
            ("Distribution of growth variable", m.distribution),
            ("lambda1 (mean AMS at 60 min)", f"{_f(m.lambda1, 2)} mm"),
            ("theta", f"{_f(m.theta, 1)} min"),
            ("eta", _f(m.eta, 3)),
            ("cv60 (L-CV at 60 min)", _f(m.cv60, 3)),
            ("beta (L-CV duration exponent)", _f(m.beta, 3)),
            ("tau3 at 60 min", _f(m.tau3, 3)),
            ("g (tau3 change per ln-duration)", _f(m.tau3_slope, 3)),
            ("RMSE of ln(mean) fit", _f(m.rmse_log_l1, 3)),
            ("RMSE of ln(L-CV) fit", _f(m.rmse_log_l2, 3)),
            ("RMSE of tau3 fit", _f(m.rmse_tau3, 3)),
            ("Notes", "; ".join(m.notes) or "-"),
        ]
        _table(doc, ["Parameter", "Value"], prow)
        with tempfile.TemporaryDirectory() as td:
            png = Path(td) / "lmom.png"
            if charts.write_lmoment_smoothing_chart(result, png):
                doc.add_picture(str(png), width=Cm(16.5))
                _caption(
                    doc,
                    "Figure 1. Sample L-moments per duration (points) and the "
                    "smooth DDF-model curves (lines).",
                )

    # 5 Recommended DDF
    doc.add_heading("5. Recommended DDF", level=1)
    if result.ddf_rows:
        durs = sorted({r["duration_min"] for r in result.ddf_rows})
        look = {(r["duration_min"], r["return_period_yr"]): r for r in result.ddf_rows}
        has_b = any(np.isfinite(r["lower_mm"]) for r in result.ddf_rows)
        header = ["Duration"] + [f"{t:g}-yr" for t in rps]
        rows = []
        for d in durs:
            row = [look[(d, rps[0])]["duration"]]
            for t in rps:
                r = look[(d, t)]
                cell = _f(r["depth_mm"], 1)
                if has_b:
                    cell += f"\n({_f(r['lower_mm'], 1)}-{_f(r['upper_mm'], 1)})"
                row.append(cell)
            rows.append(row)
        doc.add_paragraph(
            "Design depths (mm) from the duration-consistent model"
            + (", with 5-95% bootstrap bounds in brackets." if has_b else ".")
        )
        _table(doc, header, rows, font_pt=7 if len(rps) > 7 else 8)
        with tempfile.TemporaryDirectory() as td:
            png = Path(td) / "ddf.png"
            if charts.write_ddf_chart(result, png):
                doc.add_picture(str(png), width=Cm(15.5))
                _caption(
                    doc,
                    "Figure 2. Recommended DDF curves with 5-95% bands (lines) and "
                    "independent per-duration GEV quantiles (circles).",
                )
    else:
        doc.add_paragraph(
            "The duration-consistent model could not be fitted (see warnings); the "
            "recommended DDF CSV falls back to the per-duration best-fit table."
        )

    # 6 Diagnostics
    doc.add_heading("6. Diagnostics", level=1)
    doc.add_heading("6.1 Annual maximum series", level=2)
    rows = []
    for d in ds_list:
        v = np.asarray(d.values_mm, float)
        if not len(v):
            continue
        imax = int(np.argmax(v))
        rows.append(
            (
                d.duration_label,
                len(v),
                _f(v.mean(), 1),
                _f(v.min(), 1),
                _f(v.max(), 1),
                d.years[imax],
                _f(d.fixed_interval_factor, 3),
            )
        )
    _table(
        doc,
        [
            "Duration",
            "Years",
            "Mean (mm)",
            "Min (mm)",
            "Max (mm)",
            "Year of max",
            "Weiss",
        ],
        rows,
        font_pt=8,
    )
    doc.add_heading("6.2 Sample L-moments and ratio-diagram ranking", level=2)
    recs = {r.duration_label: r for r in result.recommendations}
    rows = []
    for d in ds_list:
        f = next((x for x in result.fits if x.duration_label == d.duration_label), None)
        if f is None:
            continue
        lm = f.l_moments
        rec = recs.get(d.duration_label)
        rows.append(
            (
                d.duration_label,
                _f(lm.get("l1"), 1),
                _f(lm["l2"] / lm["l1"] if lm.get("l1") else None, 3),
                _f(lm.get("t3"), 3),
                _f(lm.get("t4"), 3),
                rec.recommended_distribution if rec else "-",
                (
                    "; ".join(f"{n} {_f(dist, 3)}" for n, dist in rec.ranking[1:3])
                    if rec
                    else "-"
                ),
            )
        )
    _table(
        doc,
        ["Duration", "l1 (mm)", "L-CV", "tau3", "tau4", "Closest", "Next (distance)"],
        rows,
        font_pt=7,
    )
    with tempfile.TemporaryDirectory() as td:
        png = Path(td) / "ratio.png"
        if charts.write_ratio_diagram(result, png):
            doc.add_picture(str(png), width=Cm(13.5))
            _caption(
                doc,
                "Figure 3. L-moment ratio diagram. Scatter between durations about the "
                "theoretical curves is expected from sampling variability; a "
                "systematic position supports the chosen model distribution.",
            )

    if result.model_comparison:
        doc.add_heading("6.3 Model versus independent GEV fits", level=2)
        doc.add_paragraph(
            "Difference between the recommended (model) depth and a GEV fitted to "
            "that duration alone. Moderate differences are expected - the model "
            "pools information across durations - and are usually well within the "
            "bootstrap bounds. Large differences at one duration can indicate data "
            "problems (e.g. a single dominant storm) or that the smooth duration "
            "trend does not suit that duration."
        )
        show = [t for t in rps if round(t) in (2, 10, 50, 100)] or rps[:4]
        cl = {
            (c["duration_min"], c["return_period_yr"]): c
            for c in result.model_comparison
        }
        durs = sorted({c["duration_min"] for c in result.model_comparison})
        header = ["Duration"] + [f"{t:g}-yr model / GEV (diff)" for t in show]
        rows = []
        for d in durs:
            row = [cl[(d, show[0])]["duration"]]
            for t in show:
                c = cl.get((d, t))
                row.append(
                    f"{_f(c['model_mm'], 1)} / {_f(c['independent_gev_mm'], 1)} "
                    f"({_f(c['difference_pct'], 0)}%)"
                    if c
                    else "-"
                )
            rows.append(row)
        _table(doc, header, rows, font_pt=7)

    # 7 Warnings
    doc.add_heading("7. Warnings", level=1)
    if result.warnings:
        for w in result.warnings:
            doc.add_paragraph(w, style="List Bullet")
    else:
        doc.add_paragraph("None.")

    # 8 Limitations
    doc.add_heading("8. Limitations and use", level=1)
    for txt in (
        "Single-site frequency analysis. Record lengths of 20-40 years give "
        "reliable estimates up to roughly 2-3 times the record length; depths for "
        "rarer return periods are extrapolations and the bootstrap bounds widen "
        "accordingly. For design, compare with published regional DDF values (e.g. "
        "Smithers & Schulze 2003 / Design Rainfall for South Africa; ARR 2019 IFDs "
        "for Western Australia).",
        "Gridded products (ERA5, IMERG, CMORPH ...) represent grid-cell averages; "
        "their DDF is an areal, product-specific DDF and typically under-estimates "
        "point depths at short durations and rare return periods. Use the Adjust "
        "Sub-daily Rainfall to DDF tool to reconcile a gridded series with a "
        "gauge-based DDF.",
        "The bootstrap bounds reflect sampling uncertainty only.",
    ):
        doc.add_paragraph(txt, style="List Bullet")

    # Appendix: AMS
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.add_heading("Appendix A. Annual maximum series (mm)", level=1)
    years = sorted({y for d in ds_list for y in d.years})
    maps = [dict(zip(d.years, d.values_mm, strict=True)) for d in ds_list]
    chunk = 12
    for i in range(0, len(ds_list), chunk):
        part = list(range(i, min(i + chunk, len(ds_list))))
        header = ["Year"] + [ds_list[j].duration_label for j in part]
        rows = [[y] + [_f(maps[j].get(y), 1) for j in part] for y in years]
        _table(doc, header, rows, font_pt=7)
        doc.add_paragraph()

    doc.add_heading("References", level=1)
    for r in REFERENCES:
        doc.add_paragraph(r, style="List Bullet")
    doc.save(str(path))
