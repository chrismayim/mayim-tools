"""Word (.docx) report for Adjust Sub-daily Rainfall to DDF.

Requires python-docx (soft import, same pattern as the Design Rainfall
tool) so core.py and the CSV outputs never depend on it.
"""

from __future__ import annotations

import datetime as _dt
import math

REFERENCES = [
    "Guerreiro, S.B., et al. (2024). Unravelling the complex interplay between "
    "daily and sub-daily rainfall extremes in different climates. Weather and "
    "Climate Extremes.",
    "Lavers, D.A., et al. (2022). An evaluation of ERA5 precipitation for climate "
    "monitoring. Quarterly Journal of the Royal Meteorological Society, 148, "
    "3152-3165.",
    "Wu, G., Lv, P., Mao, Y. & Wang, K. (2024). ERA5 precipitation over China: "
    "better relative hourly and daily distribution than absolute values. Journal "
    "of Climate, 37(5), 1581-1596.",
    "Maraun, D. (2013). Bias correction, quantile mapping, and downscaling: "
    "revisiting the inflation issue. Journal of Climate, 26(6), 2137-2143.",
    "Reder, A., et al. (2022). Characterizing extreme values of precipitation at "
    "very high resolution: an experiment over twenty European cities. Weather and "
    "Climate Extremes.",
    "Hosking, J.R.M. & Wallis, J.R. (1997). Regional Frequency Analysis: An "
    "Approach Based on L-Moments. Cambridge University Press.",
    "Weiss, L.L. (1964). Ratio of true to fixed-interval maximum rainfall. Journal "
    "of the Hydraulics Division, ASCE, 90(1), 77-82.",
    "Smithers, J.C. & Schulze, R.E. (2003). Design Rainfall and Flood Estimation "
    "in South Africa. WRC Report K5/1060, Water Research Commission, Pretoria.",
    "Ball, J., et al. (eds) (2019). Australian Rainfall and Runoff: A Guide to "
    "Flood Estimation, Books 2 and 4. Commonwealth of Australia (Geoscience "
    "Australia).",
    "Efron, B. & Tibshirani, R.J. (1993). An Introduction to the Bootstrap. "
    "Chapman & Hall.",
]


def _f(v, nd=2):
    if v is None:
        return "-"
    try:
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return "-"
        return f"{float(v):.{nd}f}"
    except (TypeError, ValueError):
        return str(v)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = (
        "Light Grid Accent 1"
        if "Light Grid Accent 1" in [s.name for s in doc.styles]
        else "Table Grid"
    )
    for j, h in enumerate(header):
        t.rows[0].cells[j].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            cells[j].text = str(v)
    return t


def write_docx(result, path, chart_path=None, inputs=None):
    from docx import Document
    from docx.shared import Cm, Pt

    inputs = inputs or {}
    md = result.metadata
    ap = result.adjust_params
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    doc.add_heading("Adjustment of Sub-daily Rainfall to a Reference DDF", level=0)
    p = doc.add_paragraph()
    p.add_run(
        f"Site: {inputs.get('site') or md.get('reference_site') or '-'}    "
        f"Generated: {_dt.date.today().isoformat()}    "
        f"Tool: Mayim Tools - Adjust Sub-daily Rainfall to DDF"
    )
    s = doc.add_paragraph()
    run = s.add_run(f"Reference: {md.get('reference_tier')}. {result.status_stamp}.")
    run.bold = "INDICATIVE" in result.status_stamp

    # 1 Purpose
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "Gridded reanalysis and satellite rainfall products (e.g. ERA5) represent "
        "grid-cell averages and, in the case of reanalyses, parameterised convection. "
        "Their short-duration (sub-daily) intensities are therefore typically lower "
        "than point (gauge) intensities, and the shortfall tends to grow towards "
        "rarer events. This report documents the adjustment of the input sub-daily "
        "series so that its annual-maximum depths at all tested durations are "
        "consistent with the supplied reference Depth-Duration-Frequency (DDF) "
        "table, while retaining the timing and sequence of the recorded storms."
    )

    # 2 Inputs
    doc.add_heading("2. Inputs", level=1)
    rows = [
        ("Rainfall series", inputs.get("series", "-")),
        ("Record analysed", f"{md.get('record_start')} to {md.get('record_end')}"),
        ("Native time step", f"{_f(md.get('native_step_min'), 0)} min"),
        ("Complete years used", str(md.get("n_years_used"))),
        ("Years excluded (completeness)", md.get("years_excluded") or "none"),
        ("Time-zone offset applied", f"{_f(md.get('timezone_offset_h'), 1)} h"),
        ("Reference DDF", inputs.get("ddf", "-")),
        ("Reference tier", md.get("reference_tier")),
        ("Reference already areal (ARF applied)", str(md.get("reference_is_areal"))),
        ("Reference durations (min)", md.get("reference_durations_min")),
        ("Reference return periods (yr)", md.get("reference_return_periods_yr")),
        (
            "Reference conversions",
            "; ".join(
                x
                for x in (
                    md.get("reference_fixed_day_conversions"),
                    (
                        f"dropped: {md.get('reference_durations_dropped')}"
                        if md.get("reference_durations_dropped")
                        else ""
                    ),
                )
                if x
            )
            or "none",
        ),
        ("Distribution / fitting", md.get("distribution")),
        ("Fixed-interval correction", md.get("fixed_interval_correction")),
        (
            "Storm separation",
            f"IETD {_f(md.get('ietd_hours'), 1)} h, minimum depth "
            f"{_f(md.get('min_event_depth_mm'), 1)} mm; {md.get('n_events')} storms",
        ),
    ]
    _table(doc, ["Item", "Value"], rows)

    # 3 Method
    doc.add_heading("3. Method", level=1)
    doc.add_paragraph(
        "Annual-maximum series (AMS) are extracted from the series with sliding "
        "windows at each duration and fitted with the distribution above; return "
        "periods are annual-maximum values (AEP = 1/T), consistent with the "
        "reference table. Maxima from fixed-interval data are multiplied by the "
        "Weiss (1964) factor 1/(1 - 1/(8n)) for an n-step window. Missing data are "
        "never treated as zero; years below the completeness threshold are "
        "excluded from every duration."
    )
    mi = result.map_info or {}
    doc.add_paragraph(
        "Pass 1 - magnitude and annual total. The product's own AMS distribution G "
        "is fitted at the anchor duration "
        f"({_f(md.get('anchor_duration_min', 0) / 60, 0)} h). "
        "Each storm's anchor-duration depth x is mapped to a target depth f(x) "
        "and the whole storm is multiplied by f(x)/x, which retains storm timing. For "
        "storms in the annual-maximum range (x >= x0, where x0 is the product's "
        f"depth at T0 = {_f(mi.get('transfer_t0_yr'), 1)} yr) f(x) is the reference "
        "depth at the storm's own return period in G, T = G^-1(x): the reference table "
        "(interpolated in ln T, extrapolated beyond its rarest return period and "
        "flagged), extended between T0 and the table's most frequent return period by "
        "a GEV fitted to the table. For smaller storms f(x) = f(x0)(x/x0)^gamma, and "
        "rain outside any storm takes the factor of the smallest storm. The exponent "
        "gamma is solved so that the adjusted series reproduces the target mean "
        "annual precipitation (MAP). gamma > 1 lowers light rain relative to heavy "
        "rain, concentrating the same annual total into fewer, heavier storms; this "
        "corrects the well-documented tendency of reanalysis products to spread "
        "rainfall as too-frequent light rain with too-weak peaks (Lavers et al. 2022; "
        "Wu et al. 2024). Because pass 2 conserves mass, the MAP is retained exactly."
    )
    doc.add_paragraph(
        "Pass 2 - within-storm shape (nested peak-window scaling). Working down a "
        "ladder of calibration durations, the highest-depth window of each rung's "
        "duration inside the storm's peak window at the previous rung is multiplied "
        "by alpha = exp(a + c(ln T - ln 10)) and the remainder of that parent window "
        "by beta = (P - alpha C)/(P - C). The parent total P is unchanged, so rainfall "
        "mass is conserved at every coarser duration, zero-rain steps remain zero and "
        "storm timing is preserved. The coefficients (a, c) of each rung are fitted "
        "so that the adjusted series' implied DDF matches the reference at that "
        "rung's duration. Durations between the rungs and longer than the anchor "
        "are not fitted and serve as independent validation."
    )
    if ap.get("ensemble_n_realisations"):
        doc.add_paragraph(
            f"Stochastic ensemble. {ap.get('ensemble_n_realisations')} realisations "
            "were generated by multiplying each storm's alpha at every rung by an "
            "independent mean-one lognormal factor (sigma = "
            f"{_f(ap.get('ensemble_sigma'), 2)}), with the coefficients re-fitted "
            "so that the ensemble-median implied DDF matches the reference. The "
            "spread across realisations expresses within-storm variability that a "
            "DDF table cannot itself constrain."
        )
    doc.add_paragraph(
        f"Uncertainty. The 5-95% intervals are from {md.get('n_bootstrap')} "
        "bootstrap resamples of whole years (Efron & Tibshirani 1993). A duration is "
        "'consistent' when the median ratio is within "
        f"±{_f(md.get('tolerance_pct'), 0)}% "
        "of 1, 'inconclusive' when outside that but the interval still includes 1, "
        "and 'flatter'/'peakier' otherwise."
    )

    # 4 Results
    doc.add_heading("4. Results", level=1)
    doc.add_paragraph(
        f"Before adjustment: {md.get('overall_verdict')}. After adjustment: "
        f"{result.verdict_after or 'not adjusted'}. Validation durations only: "
        f"{ap.get('verdict_after_validation_durations', '-')}. Anchor self-check "
        f"(pass 1): {'passed' if result.self_check_passed else 'WARNING'} (max "
        f"deviation {_f(100 * md.get('self_check_max_deviation', 0), 1)}%)."
    )
    if mi:
        doc.add_paragraph("Table A. Mean annual precipitation and rainfall frequency.")
        _table(
            doc,
            ["Item", "Value"],
            [
                (
                    "MAP of input series (complete years)",
                    f"{_f(mi.get('map_input_mm'), 1)} mm",
                ),
                (
                    "Target MAP",
                    f"{_f(mi.get('map_target_mm'), 1)} mm "
                    f"(source: {mi.get('map_target_source', '-')})",
                ),
                ("MAP of adjusted series", f"{_f(mi.get('map_output_mm'), 1)} mm"),
                ("Bulk exponent gamma", _f(mi.get("transfer_gamma"), 3)),
                (
                    "Transfer start x0 (product depth at T0)",
                    f"{_f(mi.get('transfer_x0_mm'), 1)} mm",
                ),
                (
                    "Wet steps per year (> "
                    f"{_f(mi.get('wet_step_threshold_mm'), 2)} mm), input -> adjusted",
                    f"{_f(mi.get('wet_steps_per_year_input'), 0)} -> "
                    f"{_f(mi.get('wet_steps_per_year_adjusted'), 0)}",
                ),
            ],
        )
    if result.monthly_rows:
        names = [
            "Jan",
            "Feb",
            "Mar",
            "Apr",
            "May",
            "Jun",
            "Jul",
            "Aug",
            "Sep",
            "Oct",
            "Nov",
            "Dec",
        ]
        doc.add_paragraph(
            "Table B. Mean monthly rainfall (mm), input -> adjusted (seasonal check)."
        )
        _table(
            doc,
            ["Month", "Input", "Adjusted", "Ratio"],
            [
                (
                    names[r["month"] - 1],
                    _f(r["input_mean_mm"], 1),
                    _f(r["adjusted_mean_mm"], 1),
                    _f(r["ratio"], 2),
                )
                for r in result.monthly_rows
            ],
        )
    if chart_path:
        try:
            doc.add_picture(str(chart_path), width=Cm(16.5))
            cap = doc.add_paragraph().add_run(
                "Figure 1. Ratio F of implied to reference DDF depth against "
                "duration, before (left) and after (right) adjustment; bands show "
                "5-95% bootstrap intervals, the grey band the tolerance."
            )
            cap.italic = True
        except Exception:  # a missing/corrupt chart must not fail the report
            pass

    aris = sorted({r["return_period_yr"] for r in result.after_rows})
    durs = sorted({r["duration_min"] for r in result.after_rows})
    look = {(r["duration_min"], r["return_period_yr"]): r for r in result.after_rows}
    header = ["Duration", "Role"] + [f"1:{t:g} yr" for t in aris]
    rows = []
    for d in durs:
        role = {"calibration": "fitted", "validation": "check"}.get(
            look[(d, aris[0])]["role"], look[(d, aris[0])]["role"]
        )
        rows.append(
            [look[(d, aris[0])]["duration"], role]
            + [
                f"{_f(look[(d, t)]['F_before'])} -> {_f(look[(d, t)]['F_after'])}"
                for t in aris
            ]
        )
    doc.add_paragraph(
        "Table 1. Ratio F (implied / reference), before -> after. Role: fitted = "
        "calibration rung, check = independent validation duration."
    )
    _table(doc, header, rows)

    rows = []
    for d in durs:
        rows.append(
            [look[(d, aris[0])]["duration"]]
            + [
                f"{_f(look[(d, t)]['reference_depth_mm'], 1)} / "
                f"{_f(look[(d, t)]['adjusted_implied_depth_mm'], 1)}"
                for t in aris
            ]
        )
    doc.add_paragraph("Table 2. Depth (mm): reference / adjusted series.")
    _table(doc, ["Duration"] + [f"1:{t:g} yr" for t in aris], rows)

    rung_rows = []
    for k in sorted(ap):
        if k.startswith("rung_") and k.endswith("_a"):
            name = k[len("rung_") : -2]
            rung_rows.append(
                [
                    name,
                    _f(ap[k], 3),
                    _f(ap.get(f"rung_{name}_c"), 3),
                    _f(ap.get(f"rung_{name}_rmse_log"), 4),
                ]
            )
    if rung_rows:
        doc.add_paragraph("Table 3. Fitted pass-2 coefficients per calibration rung.")
        _table(doc, ["Rung", "a", "c", "RMSE (log)"], rung_rows)

    if result.ensemble_rows:
        erows = []
        el = {
            (r["duration_min"], r["return_period_yr"]): r for r in result.ensemble_rows
        }
        for d in durs:
            erows.append(
                [look[(d, aris[0])]["duration"]]
                + [
                    f"{_f(el[(d, t)]['F_ensemble_p05'])}-"
                    f"{_f(el[(d, t)]['F_ensemble_p95'])}"
                    for t in aris
                ]
            )
        doc.add_paragraph("Table 4. Ensemble spread of F (5-95% across realisations).")
        _table(doc, ["Duration"] + [f"1:{t:g} yr" for t in aris], erows)

    # 5 Diagnostics
    doc.add_heading("5. Diagnostics", level=1)
    diag = [
        ("Storms identified", md.get("n_events")),
        (
            "Storms touching missing data (not reshaped)",
            md.get("n_events_touching_missing"),
        ),
        (
            "Storms in the bulk segment (below T0)",
            md.get("n_events_bulk_segment"),
        ),
        (
            "Storms beyond reference range (extrapolated)",
            md.get("n_events_beyond_reference_range"),
        ),
        ("Storms reshaped in pass 2", ap.get("n_storms_adjusted")),
        ("Calibration durations (min)", ap.get("calibration_durations_min")),
        ("Validation durations (min)", ap.get("validation_durations_min")),
        ("Random seed", md.get("random_seed")),
    ]
    _table(doc, ["Item", "Value"], [(a, "-" if b is None else b) for a, b in diag])
    if result.warnings:
        doc.add_paragraph("Warnings:")
        for w in result.warnings:
            doc.add_paragraph(w, style="List Bullet")

    # 6 Limitations
    doc.add_heading("6. Interpretation and limitations", level=1)
    for txt in (
        "The adjusted series is a statistical reconstruction for design and "
        "continuous-simulation use. It reproduces the reference DDF in the "
        "annual-maximum sense; individual storms are not claimed to be the "
        "true historical point rainfall.",
        "No detail finer than the input's native time step is created.",
        "The target MAP should describe the same period as the series. A long-term "
        "MAP from a different period (e.g. the gauge period behind a published DDF) "
        "embeds any trend between the two periods.",
        "The reference DDF is taken as correct. Its own uncertainty (e.g. the "
        "Design Rainfall lower/upper bounds, reported in the consistency CSV) is "
        "not propagated into the adjustment.",
        "Part of the pre-adjustment deficit is the legitimate point-versus-area "
        "difference of a grid cell. If the adjusted series is to represent a "
        "catchment rather than a point, supply an areally reduced reference DDF.",
        "Validation durations are an internal consistency check; independent "
        "gauge records, where available, remain the preferred verification.",
    ):
        doc.add_paragraph(txt, style="List Bullet")

    doc.add_heading("References", level=1)
    for r in REFERENCES:
        doc.add_paragraph(r, style="List Bullet")

    doc.save(str(path))
    return True
